# -*- coding: utf-8 -*-
"""总成本表页面与导出视图（独立模块，避免改动 views.py）。"""
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from settlement.models import AccountingPeriod
from settlement.services.total_cost import (
    BLOCKS,
    SUBJECTS,
    WORKSHOP_LABELS as WS_LABELS,
    build_total_cost_data,
)


def total_cost_view(request, period_id):
    """「总成本表」已并入「生产成本」页（用户 2026-09-21）。

    这里保留 URL 只做跳转，保证旧链接/书签照旧可用（查询参数原样带过去）。
    """
    period = get_object_or_404(AccountingPeriod, pk=period_id)
    return redirect_to_production_cost(request, period)


def redirect_to_production_cost(request, period):
    from django.urls import reverse

    url = reverse("production_cost_summary", args=[period.id])
    query = request.GET.urlencode()
    return redirect("%s?%s" % (url, query) if query else url)


# ---------------- 导出 Excel ----------------

def _write_total_cost_sheet(ws, data, title="总 成 本 表"):
    """把横向宽表写进给定工作表：行=成本项目，列=四大块（期初/本期投入/产成品/在产品，各车间+合计）。"""
    period = data["period"]
    workshops = data["workshops"]
    n_ws = len(workshops)
    block_w = n_ws + 1
    ncols = 1 + len(BLOCKS) * block_w
    last_col = get_column_letter(ncols)

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    total_fill = PatternFill("solid", fgColor="FCE4D6")
    title_font = Font(name="微软雅黑", size=14, bold=True)
    normal_font = Font(name="微软雅黑", size=10)
    bold_font = Font(name="微软雅黑", size=10, bold=True)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    numfmt = "#,##0.00"

    # 标题
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"期间：{period.month:%Y年%m月}"
    ws["A2"].font = normal_font
    ws["A2"].alignment = center

    # 表头两行
    ws.merge_cells("A3:A4")
    c = ws["A3"]
    c.value = "成本项目"
    c.font = bold_font
    c.fill = header_fill
    c.border = border
    c.alignment = center

    col = 2
    for block in BLOCKS:
        l1 = get_column_letter(col)
        l2 = get_column_letter(col + block_w - 1)
        ws.merge_cells(f"{l1}3:{l2}3")
        c = ws[f"{l1}3"]
        c.value = block
        c.font = bold_font
        c.fill = header_fill
        c.border = border
        c.alignment = center
        for wi, ws_info in enumerate(workshops):
            cell = ws.cell(row=4, column=col + wi, value=ws_info["label"])
            cell.font = bold_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = center
        cell = ws.cell(row=4, column=col + n_ws, value="合计")
        cell.font = bold_font
        cell.fill = total_fill
        cell.border = border
        cell.alignment = center
        col += block_w
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 22

    def _write_row(r, values, fill=None, font=None):
        for i, v in enumerate(values):
            c2 = ws.cell(row=r, column=2 + i,
                         value="--" if v is None else float(v))
            c2.font = font or normal_font
            c2.border = border
            c2.alignment = right
            if v is not None:
                c2.number_format = numfmt
            if fill:
                c2.fill = fill

    # 数据行
    r = 5
    for row in data["rows"]:
        c = ws.cell(row=r, column=1, value=row["subject"])
        c.font = normal_font
        c.border = border
        c.alignment = Alignment(horizontal="left", vertical="center")
        _write_row(r, [cell["value"] for cell in row["cells"]])
        r += 1

    # 合计行
    c = ws.cell(row=r, column=1, value="合计")
    c.font = bold_font
    c.fill = total_fill
    c.border = border
    c.alignment = Alignment(horizontal="left", vertical="center")
    _write_row(r, [cell["value"] for cell in data["total_cells"]],
               fill=total_fill, font=bold_font)
    r += 2

    # 说明
    notes = [
        "口径说明：",
        "① 期初 = 期初在产品表按车间汇总；本期投入 = 工单结算结果 − 期初；产成品成本/在产品 = 结算结果完工/在产品拆分。",
        "② 产成品/在产品中水电/蒸汽/燃气/染料助剂/机物料/折旧按该车间制造费用结构比例拆分。",
        "③ 未结算车间本期投入/产成品/在产品无数据（--），仅有期初在产品，待结算后结转补全。",
        "④ 勾稽：期初 + 本期投入 = 产成品成本 + 在产品（对已结算车间逐科目校验，页面同步显示）。",
    ]
    for note in notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.cell(row=r, column=1, value=note).font = Font(
            name="微软雅黑", size=9, color="808080")
        ws.cell(row=r, column=1).alignment = Alignment(
            horizontal="left", vertical="center")
        r += 1

    # 列宽 / 冻结
    ws.column_dimensions["A"].width = 16
    for c2 in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(c2)].width = 10.5
    ws.freeze_panes = "B5"
    return ws


def _write_summary_sheet(ws, summary, period):
    """合并页导出的第 1 张表：车间生产成本汇总（每车间一行 + 合计行）。"""
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    total_fill = PatternFill("solid", fgColor="FCE4D6")
    title_font = Font(name="微软雅黑", size=14, bold=True)
    normal_font = Font(name="微软雅黑", size=10)
    bold_font = Font(name="微软雅黑", size=10, bold=True)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    numfmt = "#,##0.00"
    money_fmt = "#,##0.00"

    headers = ["车间", "结转状态", "工单数", "期初在制", "本期材料", "本期费用",
               "本期生产投入", "完工成本", "期末在制", "闭合差额"]
    ncols = len(headers)
    last_col = get_column_letter(ncols)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = "本月生产成本与费用"
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = (
        f"期间：{period.month:%Y年%m月}"
        f"　查看范围：{summary['scope_label']}"
    )
    ws["A2"].font = normal_font
    ws["A2"].alignment = center

    for i, name in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=name)
        c.font = bold_font
        c.fill = header_fill
        c.border = border
        c.alignment = center
    ws.row_dimensions[3].height = 22

    def _money_cell(row, col, value, font=None, fill=None, text_if_none="--"):
        # 只把数字写进单元格（合计行第 2 列是空字符串，直接 float 会报错）
        if isinstance(value, (int, float)) or hasattr(value, "quantize"):
            raw = float(value)
            number_format = money_fmt
        elif value is None:
            raw = text_if_none
            number_format = None
        else:
            raw = value
            number_format = None
        cell = ws.cell(row=row, column=col, value=raw)
        cell.font = font or normal_font
        cell.border = border
        cell.alignment = right
        if number_format:
            cell.number_format = number_format
        if fill:
            cell.fill = fill
        return cell

    r = 4
    for row in summary["rows"]:
        c = ws.cell(row=r, column=1, value=WS_LABELS.get(
            row["workshop"], row["workshop"]))
        c.font = normal_font
        c.border = border
        c.alignment = left
        for col, value in (
            (2, row["status_label"]),
            (3, row["order_count"]),
            (4, row["opening_wip"]),
            (5, row["current_material"]),
            (6, row["current_expenses"]),
            (7, row["current_input"]),
            (8, row["finished_cost"] if row["is_calculated"] else None),
            (9, row["ending_wip"] if row["is_calculated"] else None),
            (10, row["difference"] if row["is_calculated"] else None),
        ):
            if col == 2:
                cell = ws.cell(row=r, column=col, value=value)
                cell.font = normal_font
                cell.border = border
                cell.alignment = center
            elif col == 3:
                cell = ws.cell(row=r, column=col, value=value)
                cell.font = normal_font
                cell.border = border
                cell.alignment = right
            else:
                _money_cell(r, col, value)
        r += 1

    totals = summary["totals"]
    c = ws.cell(row=r, column=1, value="合计")
    c.font = bold_font
    c.fill = total_fill
    c.border = border
    c.alignment = left
    for col, value in (
        (2, ""), (3, totals["order_count"]), (4, totals["opening_wip"]),
        (5, totals["current_material"]), (6, totals["current_expenses"]),
        (7, totals["current_input"]), (8, totals["finished_cost"]),
        (9, totals["ending_wip"]), (10, totals["difference"]),
    ):
        if col == 3:
            cell = ws.cell(row=r, column=col, value=value)
            cell.font = bold_font
            cell.fill = total_fill
            cell.border = border
            cell.alignment = right
        else:
            _money_cell(r, col, value, font=bold_font, fill=total_fill,
                        text_if_none="")

    ws.column_dimensions["A"].width = 14
    for i in range(2, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.freeze_panes = "B4"
    return ws


def production_cost_export(request, period_id):
    """合并页「生产成本」的导出：一个文件两张表。

    ① 车间生产成本汇总（管理口径）
    ② 成本构成表（原总成本表，会计口径；原名「成本项目×四大块」）
    两者都跟着页面的「查看范围」车间过滤。
    """
    from settlement.services.total_cost import build_total_cost_data
    from settlement.services.workorder_costing import build_work_order_cost_summary

    from settlement.views import (
        _production_cost_target_period,
        _production_cost_period_options,
    )

    period = get_object_or_404(AccountingPeriod, pk=period_id)
    legacy = str(request.GET.get("workshop", "")).strip()
    summary_workshop = str(request.GET.get("p_workshop", "")).strip() or legacy
    matrix_workshop = str(request.GET.get("m_workshop", "")).strip() or legacy
    summary_period = _production_cost_target_period(
        period, str(request.GET.get("p_period", "")).strip()
    )
    matrix_period = _production_cost_target_period(
        period, str(request.GET.get("m_period", "")).strip()
    )
    summary = build_work_order_cost_summary(summary_period, summary_workshop)
    data = build_total_cost_data(matrix_period, only_workshop=matrix_workshop)
    period = summary_period      # 第 1 张表的表头用车间汇总那个期间
    only_workshop = summary_workshop

    wb = Workbook()
    _write_summary_sheet(wb.active, summary, period)
    wb.active.title = "车间生产成本汇总"
    matrix_ws = wb.create_sheet("成本构成表")
    _write_total_cost_sheet(matrix_ws, data, title="成本构成表")

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    suffix = f"-{only_workshop}" if only_workshop else ""
    filename = f"生产成本-{period.month:%Y-%m}{suffix}.xlsx"
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response


def total_cost_export(request, period_id):
    """旧的「总成本表」导出地址：转到合并页的导出（两个工作表），参数原样带过去。"""
    from django.urls import reverse

    url = reverse("production_cost_export", args=[period_id])
    query = request.GET.urlencode()
    return redirect("%s?%s" % (url, query) if query else url)


def production_cost_matrix_export(request, period_id):
    """「成本构成表」单张工作表导出（用户 2026-09-27；原名「成本项目 × 四大块」）。

    与页面同一套取数（build_total_cost_data），跟着该表自己的筛选
    （m_period / m_workshop，老参数 workshop 仍生效）。
    """
    period = get_object_or_404(AccountingPeriod, pk=period_id)
    legacy = str(request.GET.get("workshop", "")).strip()
    only_workshop = (
        str(request.GET.get("m_workshop", "")).strip() or legacy
    )
    from settlement.views import _production_cost_target_period

    target_period = _production_cost_target_period(
        period, str(request.GET.get("m_period", "")).strip()
    )
    data = build_total_cost_data(target_period, only_workshop=only_workshop)

    wb = Workbook()
    # 名称更换（用户 2026-09-27：原名「成本项目×四大块」太土）
    _write_total_cost_sheet(wb.active, data, title="成本构成表")
    wb.active.title = "成本构成表"

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    suffix = f"-{only_workshop}" if only_workshop else ""
    filename = f"成本构成表-{target_period.month:%Y-%m}{suffix}.xlsx"
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return response
