from __future__ import annotations

from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from settlement.models import MatchAllocation
from settlement.services.matching import is_purchase_receipt, is_sales_issue
from settlement.services.purchase_accruals import (
    build_purchase_accrual_rows,
    build_purchase_difference_rows,
)
from settlement.services.sales_uninvoiced import (
    build_goods_issued_rollforward,
    build_sales_uninvoiced_rows,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBTOTAL_FILL = PatternFill("solid", fgColor="D9EAF7")
# 数字列统一显示：千分位分隔 + 两位小数（底层仍是数值，可继续计算）
NUMBER_FORMAT = "#,##0.00"


def _is_number(value) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(
        value, bool
    )


def _clean_cell_value(value):
    """Excel 不支持带时区的日期时间：统一转成本地时间的 naive datetime。

    Django 开启 USE_TZ 后 DateTimeField 返回的 datetime 带 tzinfo，
    直接写入 openpyxl 会抛
    "Excel does not support timezones in datetimes"。
    这里在写入前做全局兜底，覆盖所有报表。
    """
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return timezone.localtime(value).replace(tzinfo=None)
        return value
    return value


def _clean_row(row):
    return [_clean_cell_value(value) for value in row]


def _append_table(workbook, title, headers, rows):
    worksheet = workbook.create_sheet(title=title)
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        worksheet.append(_clean_row(row))
    # 数据区数字单元格设置千分位+两位小数格式
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if _is_number(cell.value):
                cell.number_format = NUMBER_FORMAT
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, 1):
        width = max(
            len(str(cell.value or "")) for cell in list(column)[:200]
        )
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(width + 2, 10), 42
        )
    return worksheet


def read_uploaded_rows(path):
    """读取用户上传的原始表格文件，返回 (表头, 数据行列表)。保留原始行数与列。"""
    suffix = path.suffix.lower()
    rows = []
    header = []
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header = list(next(it))
        except StopIteration:
            wb.close()
            return [], []
        for r in it:
            if any(v not in (None, "") for v in r):
                rows.append(tuple(r))
        wb.close()
    elif suffix == ".csv":
        import csv
        import io

        raw = path.read_bytes()
        decoded = None
        for enc in ("utf-8-sig", "gb18030", "utf-8"):
            try:
                decoded = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            return [], []
        csv_rows = list(csv.reader(io.StringIO(decoded)))
        if not csv_rows:
            return [], []
        header = csv_rows[0]
        rows = [
            tuple(r)
            for r in csv_rows[1:]
            if any(str(v).strip() for v in r)
        ]
    else:
        return [], []
    # 去掉尾部空列
    while header and header[-1] in (None, ""):
        header.pop()
        rows = [r[: len(header)] for r in rows]
    return header, rows


def build_settlement_workbook(period):
    workbook = Workbook()
    workbook.remove(workbook.active)

    _append_table(
        workbook,
        "库存计价",
        [
            "物料编码",
            "物料名称",
            "期初数量",
            "期初金额",
            "本期入库数量",
            "本期入库金额",
            "月末加权平均价",
            "本期出库数量",
            "本期出库金额",
            "期末数量",
            "期末金额",
            "是否收敛",
        ],
        (
            (
                item.material_code,
                item.material_name,
                item.opening_quantity,
                item.opening_amount,
                item.receipt_quantity,
                item.receipt_amount,
                item.weighted_average_price,
                item.issue_quantity,
                item.issue_amount,
                item.ending_quantity,
                item.ending_amount,
                "是" if item.converged else "否",
            )
            for item in period.inventory_valuations.all()
        ),
    )

    carry_rows = []
    for item in period.opening_inventory_balances.all():
        carry_rows.append(
            (
                "期初库存",
                item.source_period.month if item.source_period else None,
                item.material_code,
                item.material_name,
                item.quantity,
                item.amount,
                item.unit_price,
                (
                    "系统自动承接"
                    if item.is_system_generated
                    else "人工上传/覆盖"
                ),
            )
        )
    for item in period.opening_wip.all():
        carry_rows.append(
            (
                "期初在产品",
                item.source_period.month if item.source_period else None,
                item.order_no,
                "",
                "",
                item.total,
                "",
                "系统自动承接" if item.is_system_generated else "人工上传/覆盖",
            )
        )
    for item in period.purchase_contracts.filter(is_system_generated=True):
        carry_rows.append(
            (
                "合同余额",
                item.carried_from.period.month if item.carried_from else None,
                item.contract_no,
                item.material_name,
                item.quantity,
                "",
                item.unit_price_excl_tax,
                item.supplier_name,
            )
        )
    _append_table(
        workbook,
        "跨期结转",
        [
            "结转类型",
            "来源期间",
            "编码/单号",
            "名称",
            "数量",
            "金额",
            "单价",
            "说明",
        ],
        carry_rows,
    )

    _append_table(
        workbook,
        "期初采购暂估",
        [
            "原入库期间",
            "原入库日期",
            "原入库单号",
            "行号",
            "供应商",
            "合同号",
            "物料编码",
            "物料名称",
            "原入库数量",
            "以前到票数量",
            "期初未到票数量",
            "暂估单价",
            "期初暂估金额",
            "剩余未到票数量",
            "最近到票期间",
            "匹配状态",
        ],
        (
            (
                item.original_period,
                item.document_date,
                item.document_no,
                item.line_no,
                item.supplier_name or item.supplier_code,
                item.contract_no,
                item.material_code,
                item.material_name,
                item.original_quantity,
                item.prior_invoiced_quantity,
                item.opening_quantity,
                item.provisional_unit_price,
                item.opening_amount,
                item.remaining_quantity,
                item.invoice_period,
                item.get_match_status_display(),
            )
            for item in period.opening_purchase_accruals.all()
        ),
    )

    _append_table(
        workbook,
        "期初发出商品",
        [
            "原出库期间",
            "原出库日期",
            "原出库单号",
            "行号",
            "销售订单号",
            "客户",
            "产品编码",
            "产品名称",
            "原出库数量",
            "以前开票数量",
            "期初未开票数量",
            "原出库单位成本",
            "期初发出商品成本",
            "剩余未开票数量",
            "最近开票期间",
            "匹配状态",
        ],
        (
            (
                item.original_period,
                item.document_date,
                item.document_no,
                item.line_no,
                item.sales_order_no,
                item.customer_name or item.customer_code,
                item.material_code,
                item.material_name,
                item.original_quantity,
                item.prior_invoiced_quantity,
                item.opening_quantity,
                item.unit_cost,
                item.opening_amount,
                item.remaining_quantity,
                item.invoice_period,
                item.get_match_status_display(),
            )
            for item in period.opening_goods_issued.all()
        ),
    )

    _append_table(
        workbook,
        "期末存货",
        [
            "期末日期",
            "物料号",
            "物料名称",
            "分类名称",
            "仓库名称",
            "批次号",
            "结存数量",
            "单位名称",
            "单价",
            "金额",
        ],
        (
            (
                item.ending_date,
                item.material_code,
                item.material_name,
                item.category_name,
                item.warehouse,
                item.batch_no,
                item.quantity,
                item.uom,
                item.unit_price,
                item.amount,
            )
            for item in period.ending_inventory_balances.all()
        ),
    )

    _append_table(
        workbook,
        "车间结转",
        [
            "结转顺序",
            "车间/成本中心",
            "状态",
            "工单成本",
            "完工入库成本",
            "循环次数",
            "是否收敛",
            "结转时间",
            "结转人",
        ],
        (
            (
                item.sequence,
                item.workshop,
                item.get_status_display(),
                item.order_cost,
                item.production_receipt_amount,
                item.iterations,
                "是" if item.converged else "否",
                (
                    timezone.localtime(item.settled_at).replace(tzinfo=None)
                    if item.settled_at
                    else None
                ),
                str(item.settled_by or ""),
            )
            for item in period.workshop_settlements.all()
        ),
    )

    _append_table(
        workbook,
        "库存流水",
        [
            "创建时间",
            "物料名称",
            "物料编号",
            "分类",
            "物料规格",
            "仓库名称",
            "批号",
            "操作数量",
            "主单位",
            "辅助数量",
            "辅助单位",
            "备注",
            "业务单据编号",
            "业务类型",
            "供应商名称",
            "单价",
            "金额",
        ],
        (
            (
                item.txn_date,
                item.material_name,
                item.material_code,
                item.category_name,
                item.specification,
                item.warehouse,
                item.batch_no,
                item.quantity,
                item.uom,
                item.auxiliary_quantity,
                item.auxiliary_uom,
                item.remark,
                item.document_no,
                item.business_type,
                item.supplier_name or item.customer_name
                or item.supplier_code or item.customer_code or "",
                item.unit_price,
                item.amount,
            )
            for item in period.inventory_transactions.all()
        ),
    )

    purchase_rows = []
    for txn in period.inventory_transactions.all():
        if not is_purchase_receipt(txn):
            continue
        contract_allocs = txn.allocations.filter(
            kind=MatchAllocation.Kind.PURCHASE_CONTRACT
        ).select_related("contract_line")
        invoice_allocs = txn.allocations.filter(
            kind=MatchAllocation.Kind.PURCHASE_INVOICE,
            period__month__lte=period.month,
        ).select_related("purchase_invoice_line")
        purchase_rows.append(
            (
                txn.id,
                txn.txn_date,
                txn.document_no,
                txn.supplier_name,
                txn.material_code,
                txn.material_name,
                txn.quantity,
                "；".join(
                    f"{item.contract_line.contract_no}:{item.matched_quantity}"
                    for item in contract_allocs
                ),
                "；".join(
                    f"{item.purchase_invoice_line.invoice_no}:{item.matched_quantity}"
                    for item in invoice_allocs
                ),
                txn.unit_price,
                txn.amount,
                txn.price_source,
                txn.invoice_period,
            )
        )
    _append_table(
        workbook,
        "采购匹配",
        [
            "库存流水ID",
            "入库日期",
            "入库单号",
            "供应商",
            "物料编码",
            "物料名称",
            "入库数量",
            "匹配合同及数量",
            "匹配发票及数量",
            "暂估入库单价",
            "固定暂估入库金额",
            "价格来源",
            "最新发票期间",
        ],
        purchase_rows,
    )

    accrual_rows = build_purchase_accrual_rows(period)
    _append_table(
        workbook,
        "月末采购暂估",
        [
            "入库期间",
            "入库日期",
            "入库单号",
            "合同号",
            "供应商",
            "物料编码",
            "物料名称",
            "入库数量",
            "累计到票数量",
            "尚未到票数量",
            "暂估单价",
            "月末暂估金额",
            "最近到票期间",
            "暂估依据",
            "状态",
        ],
        (
            (
                row.receipt_month,
                row.document_date,
                row.document_no,
                row.contract_nos,
                row.supplier_name or row.supplier_code,
                row.material_code,
                row.material_name,
                row.receipt_quantity,
                row.invoiced_quantity,
                row.outstanding_quantity,
                row.provisional_unit_price,
                row.provisional_amount,
                row.latest_invoice_period,
                row.price_source,
                row.status,
            )
            for row in accrual_rows
        ),
    )

    difference_rows = build_purchase_difference_rows(period)
    _append_table(
        workbook,
        "采购暂估差异",
        [
            "发票期间",
            "发票号",
            "供应商",
            "物料编码",
            "物料名称",
            "入库期间",
            "入库日期",
            "入库单号",
            "核销数量",
            "暂估单价",
            "暂估金额",
            "发票单价",
            "发票金额",
            "暂估价差",
            "是否跨期",
            "状态",
        ],
        (
            (
                row.invoice.invoice_period,
                row.invoice.invoice_no,
                row.invoice.supplier_name or row.invoice.supplier_code,
                row.material_code,
                row.material_name,
                row.receipt_month,
                row.document_date,
                row.document_no,
                row.matched_quantity,
                row.provisional_unit_price,
                row.provisional_amount,
                row.invoice_unit_price,
                row.invoice_amount,
                row.difference_amount,
                "是" if row.cross_period else "否",
                row.status,
            )
            for row in difference_rows
        ),
    )

    sales_rows = []
    for txn in period.inventory_transactions.all():
        if not is_sales_issue(txn):
            continue
        allocations = txn.allocations.filter(
            kind=MatchAllocation.Kind.SALES_INVOICE
        ).select_related("sales_invoice_line")
        sales_rows.append(
            (
                txn.id,
                txn.txn_date,
                txn.document_no,
                txn.customer_name,
                txn.material_code,
                txn.material_name,
                txn.quantity,
                txn.amount,
                "；".join(
                    f"{item.sales_invoice_line.invoice_no}:{item.matched_quantity}"
                    for item in allocations
                ),
                txn.invoice_period,
                txn.get_match_status_display(),
            )
        )
    _append_table(
        workbook,
        "销售匹配",
        [
            "库存流水ID",
            "出库日期",
            "出库单号",
            "客户",
            "产品编码",
            "产品名称",
            "出库数量",
            "出库成本",
            "匹配发票及数量",
            "开票期间",
            "匹配状态",
        ],
        sales_rows,
    )

    uninvoiced_rows = build_sales_uninvoiced_rows(period)
    _append_table(
        workbook,
        "月末已出库未开票",
        [
            "出库期间",
            "出库日期",
            "出库单号",
            "销售订单号",
            "客户",
            "产品编码",
            "产品名称",
            "出库数量",
            "累计开票数量",
            "未开票数量",
            "出库成本",
            "已开票成本",
            "未开票成本",
            "匹配发票",
            "最近开票期间",
            "未开票天数",
            "状态",
        ],
        (
            (
                row.issue_month,
                row.document_date,
                row.document_no,
                row.sales_order_no,
                row.customer_name or row.customer_code,
                row.material_code,
                row.material_name,
                row.issue_quantity,
                row.invoiced_quantity,
                row.outstanding_quantity,
                row.issue_amount,
                row.invoiced_cost,
                row.outstanding_cost,
                row.invoice_nos,
                row.latest_invoice_period,
                row.outstanding_days,
                row.status,
            )
            for row in uninvoiced_rows
        ),
    )

    goods_issued = build_goods_issued_rollforward(period)
    _append_table(
        workbook,
        "发出商品滚动",
        [
            "核算期间",
            "期初发出商品",
            "本月销售出库",
            "本月开票结转成本",
            "期末发出商品",
            "差额",
            "期初未开票笔数",
            "期末未开票笔数",
            "其中以前月份出库",
        ],
        [
            (
                period.month,
                goods_issued["opening_amount"],
                goods_issued["current_issue_amount"],
                goods_issued["invoice_cleared_amount"],
                goods_issued["ending_amount"],
                goods_issued["difference"],
                goods_issued["opening_count"],
                goods_issued["ending_count"],
                goods_issued["prior_period_count"],
            )
        ],
    )

    _append_table(
        workbook,
        "工单成本",
        [
            "工单编号",
            "车间",
            "产品编码",
            "产品名称",
            "计划数量",
            "报工数量",
            "在产品数量",
            "材料成本",
            "水电",
            "蒸汽",
            "天然气",
            "折旧",
            "机物料",
            "染料助剂",
            "工资",
            "工单总成本",
            "完工成本",
            "完工单位成本",
            "在产品成本",
        ],
        (
            (
                item.work_order.order_no,
                item.work_order.production_dept_code,
                item.product_code,
                item.product_name,
                item.plan_quantity,
                item.finished_quantity,
                item.wip_quantity,
                item.direct_material,
                item.water_power,
                item.steam,
                item.natural_gas,
                item.depreciation,
                item.consumables,
                item.dye_auxiliary,
                item.wages,
                item.total_cost,
                item.finished_cost,
                item.finished_unit_cost,
                item.wip_cost,
            )
            for item in period.work_order_cost_results.select_related(
                "work_order"
            ).all()
        ),
    )

    _append_table(
        workbook,
        "勾稽结果",
        [
            "检查编码",
            "检查项目",
            "左侧金额",
            "右侧金额",
            "差额",
            "状态",
            "分项明细",
        ],
        (
            (
                item.check_code,
                item.check_name,
                item.left_amount,
                item.right_amount,
                item.difference,
                item.get_status_display(),
                "；".join(
                    f"{key}={value}"
                    for key, value in item.details.items()
                ),
            )
            for item in period.reconciliations.all()
        ),
    )

    _append_table(
        workbook,
        "异常清单",
        ["严重程度", "类别", "引用", "异常内容", "处理状态"],
        (
            (
                item.get_severity_display(),
                item.category,
                item.reference,
                item.message,
                item.get_status_display(),
            )
            for item in period.exceptions.all()
        ),
    )

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def settlement_excel_response(period):
    stream = build_settlement_workbook(period)
    filename = f"成本结算结果-{period.month:%Y%m}.xlsx"
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f"attachment; filename*=UTF-8''{__import__('urllib.parse').parse.quote(filename)}"
    )
    return response


def build_listing_workbook(sheets):
    """生成单文件多工作表 Excel。

    sheets: [(工作表名, 表头列表, 数据行(可迭代的元组/列表)), ...]
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in sheets:
        _append_table(workbook, title, headers, rows)
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream

