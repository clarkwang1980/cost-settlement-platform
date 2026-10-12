# -*- coding: utf-8 -*-
"""单价表（入库/出库 · 跨月对比）。

用户口径（2026-09-12 确认）：
- 一个工作簿两个工作表：「入库单价」与「出库单价」，两个表结构相同；
- 列固定为：序号、品类、品名编码、品名名称，加该年度 1~12 月（如 2026年1月 … 2026年12月）；
- 某个月只要有新入库/新出库，就把该月单价填到对应月份列（没有的月份留空）；
- 目的：一眼看出每个月入库价、出库价的变化。

取价口径（与平台既有口径一致）：
- 入库单价 = 该月入库金额 ÷ 入库数量。
  入库含采购入库、完工核算入库（报工/生产入库）、其他入库；
  仓库间调拨/库存调整/库存转换/报工调整属内部流转，不计入（导入时已标为忽略匹配）。
- 出库单价 = 该月出库金额 ÷ 出库数量（月末加权平均价，全月同一价）。
- 只有「当月确有入库/出库」的月份才填；没有进出的月份留空。
- 该月有进出的物料，若上传过「历史单价」则以上传价为准（与单价表页面口径一致）。
"""
from __future__ import annotations

from decimal import Decimal

ZERO = Decimal("0")

MONTH_HEADER_FMT = "{year}年{month}月"

# 品类展示顺序（其余品类按名称排在后面）
CATEGORY_ORDER = [
    "原料",
    "染色纤维",
    "纤维",
    "熟条",
    "浆粕",
    "纱线",
    "面料",
    "坯布",
    "产成品",
]

SHEET_INBOUND = "入库单价"
SHEET_OUTBOUND = "出库单价"

BASE_HEADERS = ["序号", "品类", "品名编码", "品名名称"]


def month_headers(year: int) -> list[str]:
    return [MONTH_HEADER_FMT.format(year=year, month=m) for m in range(1, 13)]


def _category_sort_key(category: str):
    try:
        return (0, CATEGORY_ORDER.index(category), "")
    except ValueError:
        return (1, 0, category or "")


def _price(amount, quantity):
    """金额 ÷ 数量；数量为 0 或空返回 None。"""
    qty = Decimal(quantity or ZERO)
    if qty == ZERO:
        return None
    return (Decimal(amount or ZERO) / qty).quantize(Decimal("0.00000001"))


def internal_transfer_qty(periods):
    """内部流转（调拨/库存调整/库存转换/报工调整）的入库/出库数量：{(月份, 编码): (in, out)}。

    用户 2026-09-30：这些行**算期末数量**（它们是真实的数量移动），但**不算成本**（没有金额），
    所以计算「入库单价＝入库金额÷入库数量」时必须把它们从分母里扣掉，否则单价被摊薄。
    """
    from collections import defaultdict

    from settlement.services.costing import is_internal_transfer

    result = defaultdict(lambda: [ZERO, ZERO])
    for period in periods:
        month = period.month.month
        for txn in period.inventory_transactions.all():
            if not is_internal_transfer(txn):
                continue
            qty = abs(Decimal(txn.quantity or ZERO))
            slot = result[(month, txn.material_code)]
            if txn.direction == "in":
                slot[0] += qty
            else:
                slot[1] += qty
    return result


def build_price_matrix(period, months=None, code_filter=""):
    """按年度汇总各物料的月度入库/出库单价。

    返回 dict(year=, inbound=[row], outbound=[row])，
    row = {seq, category, code, name, prices: {月份: Decimal|None}}。
    """
    from settlement.models import (
        AccountingPeriod,
        HistoricalPrice,
        InventoryTransaction,
        InventoryValuation,
        Material,
    )

    year = period.month.year
    periods = list(AccountingPeriod.objects.filter(month__year=year))
    month_to_period = {p.month.month: p for p in periods}
    wanted_months = sorted(months) if months else list(range(1, 13))
    wanted_months = [m for m in wanted_months if m in month_to_period] or sorted(month_to_period)
    internal_qty = internal_transfer_qty(periods)

    valuations = {}  # (月份, 编码) -> InventoryValuation
    historical = {}  # (月份, 编码) -> HistoricalPrice
    name_by_code = {}
    for p in periods:
        for v in InventoryValuation.objects.filter(period=p):
            valuations[(p.month.month, v.material_code)] = v
            if v.material_name:
                name_by_code.setdefault(v.material_code, v.material_name)
        for hp in HistoricalPrice.objects.filter(period=p):
            historical[(p.month.month, hp.material_code)] = hp
            if hp.material_name:
                name_by_code.setdefault(hp.material_code, hp.material_name)

    category_by_code = {}
    for code, name, cat in (
        InventoryTransaction.objects.filter(period__in=periods)
        .order_by("material_code")
        .values_list("material_code", "material_name", "category_name")
    ):
        if name and not name_by_code.get(code):
            name_by_code[code] = name
        if cat and code not in category_by_code:
            category_by_code[code] = cat

    for mat in Material.objects.all().only("code", "name"):
        if mat.name:
            name_by_code.setdefault(mat.code, mat.name)

    codes = {code for (_month, code) in valuations}
    codes |= {code for (_month, code) in historical}
    if code_filter:
        needle = code_filter.strip().lower()
        codes = {code for code in codes if needle in code.lower()}

    inbound_rows = []
    outbound_rows = []
    for code in codes:
        in_prices = {}
        out_prices = {}
        for month in wanted_months:
            v = valuations.get((month, code))
            hp = historical.get((month, code))
            in_price = None
            out_price = None
            # 只填「当月确有入库/出库」的月份（按你的口径：有进出的才补到对应月份）
            if v is not None:
                in_qty_internal, out_qty_internal = internal_qty.get((month, code), [ZERO, ZERO])
                in_qty = Decimal(v.receipt_quantity or ZERO) - Decimal(in_qty_internal)
                out_qty = Decimal(v.issue_quantity or ZERO) - Decimal(out_qty_internal)
                if in_qty > ZERO:
                    in_price = _price(v.receipt_amount, in_qty)
                if out_qty > ZERO:
                    out_price = _price(v.issue_amount, out_qty)
            # 该月有进出的，若上传过历史单价则以上传价为准（与单价表页面口径一致）
            if hp is not None:
                if in_price is not None and hp.receipt_price is not None:
                    in_price = Decimal(hp.receipt_price)
                if out_price is not None and hp.issue_price is not None:
                    out_price = Decimal(hp.issue_price)
            if in_price is not None:
                in_prices[month] = in_price
            if out_price is not None:
                out_prices[month] = out_price
        if in_prices:
            inbound_rows.append(
                {
                    "category": category_by_code.get(code, ""),
                    "code": code,
                    "name": name_by_code.get(code, ""),
                    "prices": in_prices,
                }
            )
        if out_prices:
            outbound_rows.append(
                {
                    "category": category_by_code.get(code, ""),
                    "code": code,
                    "name": name_by_code.get(code, ""),
                    "prices": out_prices,
                }
            )

    def _sorted(rows):
        rows.sort(key=lambda r: (_category_sort_key(r["category"]), r["code"]))
        for index, row in enumerate(rows, 1):
            row["seq"] = index
        return rows

    return {
        "year": year,
        "inbound": _sorted(inbound_rows),
        "outbound": _sorted(outbound_rows),
    }


def _write_sheet(
    ws,
    title: str,
    year: int,
    rows,
    *,
    number_format: str = "#,##0.0000",
    months=None,
):
    """写一个工作表：标题行 + 表头行 + 数据行。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    months = sorted(months) if months else list(range(1, 13))
    headers = BASE_HEADERS + [MONTH_HEADER_FMT.format(year=year, month=m) for m in months]
    ws.append([title])
    ws.merge_cells(
        start_row=1, start_column=1, end_row=1, end_column=len(headers)
    )
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.append(headers)

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="1F4E79")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(2, col)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in rows:
        line = [row["seq"], row["category"], row["code"], row["name"]]
        line += [row["prices"].get(month) for month in months]
        ws.append(line)
        current = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(current, col)
            cell.border = border
            if col > len(BASE_HEADERS):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = number_format

    widths = [6, 12, 16, 32]
    for index in range(1, len(headers) + 1):
        if index <= len(widths):
            width = widths[index - 1]
        else:
            width = 14
        ws.column_dimensions[
            ws.cell(2, index).column_letter
        ].width = width
    ws.freeze_panes = ws.cell(3, len(BASE_HEADERS) + 1).coordinate


def build_price_table_workbook(period, n=None, months=None, code_filter=""):
    """生成单价表工作簿（两个工作表：入库单价 / 出库单价）。

    用户 2026-09-30：只输出**所选期间**的月份列（默认页面筛出来的那几期），
    并支持按物料编码筛选 —— 否则永远 1~12 月、11 列空着，表看着特别大。
    """
    from io import BytesIO

    from openpyxl import Workbook

    data = build_price_matrix(period, months=months, code_filter=code_filter)
    year = data["year"]
    used_months = sorted({m for row in (data["inbound"] + data["outbound"]) for m in row["prices"]}) or (sorted(months) if months else list(range(1, 13)))

    wb = Workbook()
    ws_in = wb.active
    ws_in.title = SHEET_INBOUND
    _write_sheet(ws_in, f"入库单价表（{year}年）", year, data["inbound"], months=used_months)

    ws_out = wb.create_sheet(SHEET_OUTBOUND)
    _write_sheet(ws_out, f"出库单价表（{year}年）", year, data["outbound"], months=used_months)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def build_price_template_workbook(year=None, sample_rows=2):
    """下载用的空白模板：与单价表同样结构，附示例行（编码以「示例」开头，导入时自动跳过）。"""
    from io import BytesIO

    from openpyxl import Workbook
    from settlement.models import AccountingPeriod

    if year is None:
        latest = AccountingPeriod.objects.order_by("-month").first()
        from datetime import date

        year = latest.month.year if latest else date.today().year

    sample_in = [
        {
            "seq": index + 1,
            "category": "纱线",
            "code": f"示例-P00{index + 1}",
            "name": f"示例产品{chr(65 + index)}",
            "prices": {},
        }
        for index in range(sample_rows)
    ]
    if sample_in:
        sample_in[0]["prices"] = {7: Decimal("77.56990935")}
    sample_out = [
        {
            "seq": index + 1,
            "category": "纱线",
            "code": f"示例-P00{index + 1}",
            "name": f"示例产品{chr(65 + index)}",
            "prices": {},
        }
        for index in range(sample_rows)
    ]
    if sample_out:
        sample_out[0]["prices"] = {7: Decimal("81.23591006")}

    wb = Workbook()
    ws_in = wb.active
    ws_in.title = SHEET_INBOUND
    _write_sheet(
        ws_in,
        f"入库单价表（{year}年）—— 每月有新入库就把该月单价填到对应月份列",
        year,
        sample_in,
    )
    ws_out = wb.create_sheet(SHEET_OUTBOUND)
    _write_sheet(
        ws_out,
        f"出库单价表（{year}年）—— 每月有新出库就把该月单价填到对应月份列",
        year,
        sample_out,
    )

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
