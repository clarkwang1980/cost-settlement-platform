from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone as dj_timezone
from openpyxl import load_workbook

from settlement.models import (
    AccountingPeriod,
    BomLine,
    EndingInventoryBalance,
    HistoricalPrice,
    ImportBatch,
    InventoryTransaction,
    InventoryValuation,
    IssueProgress,
    Material,
    MaterialIssue,
    MatchAllocation,
    MatchingMode,
    OpeningGoodsIssued,
    OpeningInventoryBalance,
    OpeningPurchaseAccrual,
    OpeningWip,
    PurchaseContractLine,
    PurchaseInvoiceLine,
    PulpMaterial,
    ReconciliationResult,
    SalesInvoiceLine,
    SalesLedgerLine,
    SettlementRun,
    WorkOrder,
    WorkOrderCostResult,
    WorkOrderDetail,
    WorkReport,
    WorkshopSettlement,
    WorkshopCostPool,
)


HEADER_CLEANER = re.compile(r"[\s\r\n\t（）()【】\[\]：:、,，._\-*＊#※]+")


def normalize_header(value) -> str:
    return HEADER_CLEANER.sub("", str(value or "").strip().lower())


ALIASES = {
    "inventory": {
        "txn_date": ["创建时间", "业务日期", "单据日期", "日期", "出入库日期", "记账日期", "入库时间"],
        "document_no": ["业务单据编号", "单据号", "单据编号", "出入库单号", "凭证号"],
        "line_no": ["行号", "序号", "明细行"],
        "business_type": ["业务类型", "出入库类型", "单据类型", "事务类型"],
        "direction": ["方向", "出入库方向", "收发标志"],
        "warehouse": ["仓库名称", "仓库", "库位"],
        "material_code": ["物料编号", "物料编码", "物料代码", "存货编码", "产品编码", "品号"],
        "material_name": ["物料名称", "存货名称", "产品名称", "品名"],
        "specification": ["物料规格", "规格", "规格型号", "型号"],
        "category_name": ["分类", "分类名称", "物料分类", "类别"],
        "batch_no": ["批号", "批次号", "批次"],
        "quantity": ["操作数量", "数量", "实收数量", "实发数量", "出入库数量", "基本数量", "入库数量"],
        "auxiliary_quantity": ["辅助数量", "辅数量", "辅助计量数量"],
        "uom": ["主单位", "单位", "计量单位", "基本单位"],
        "auxiliary_uom": ["辅助单位", "辅单位"],
        "order_no": ["工单号", "生产订单号", "生产订单", "订单号"],
        "workshop": ["车间", "生产车间", "成本中心"],
        "supplier_code": ["供应商编码", "供应商代码"],
        "supplier_name": ["供应商名称", "供应商", "供应商全称"],
        "customer_code": ["客户编码", "客户代码"],
        "customer_name": ["客户", "客户名称"],
        "contract_no": ["合同号", "采购合同号", "合同编号"],
        "invoice_no": ["发票号", "发票号码"],
        "unit_price": ["单价", "不含税单价", "入库单价", "单位成本"],
        "amount": ["金额", "不含税金额", "成本金额", "入库金额"],
        "remark": ["备注", "摘要"],
    },
    "purchase_contract": {
        "contract_no": ["合同号", "采购合同号", "合同编号", "合同编码"],
        "contract_date": ["合同日期", "签订日期", "日期", "创建时间"],
        "supplier_code": ["供应商编码", "供应商代码"],
        "supplier_name": ["供应商", "供应商名称"],
        "material_code": ["物料编码", "物料代码", "存货编码", "品名编码"],
        "material_name": ["物料名称", "存货名称", "品名"],
        "quantity": ["合同数量", "数量"],
        "received_quantity": ["已签收数量", "签收数量"],
        # 用户 2026-09-19：模板最左新增「收货期间」列（合同侧只记录与展示）
        "receipt_period": ["收货期间", "收货月份", "入库期间", "入库月份"],
        # 用户 2026-09-19：紧挨「收货期间」再加一列「收票日期」（发票收到日期）
        "invoice_date": ["收票日期", "发票日期", "开票日期", "收票月份"],
        # 用户 2026-09-19：合同表的「不含税总价」，用于「合同金额」列（与上传文件口径一致）
        "amount_excl_tax": [
            "原币未税总价（报表用）",
            "原币未税总价",
            "不含税总价",
            "未税总价",
            "合同金额",
            "原币未税总金额（报表用）",
        ],
        "unit_price_excl_tax": [
            "不含税单价",
            "合同不含税价",
            "单价",
            "原币未税单价（报表用）",
            "原币未税单价",
        ],
    },
    "purchase_invoice": {
        "invoice_no": ["发票号", "发票号码"],
        "invoice_date": ["收票日期", "发票日期", "开票日期", "日期"],
        "invoice_period": ["发票期间", "开票期间", "所属期间"],
        "contract_no": ["合同号", "采购合同号", "合同编号"],
        "supplier_code": ["供应商编码", "供应商代码"],
        "supplier_name": ["供应商", "供应商名称"],
        "material_code": ["物料编码", "物料代码", "存货编码"],
        "material_name": ["物料名称", "存货名称", "品名"],
        "uom": ["单位", "计量单位"],
        "quantity": ["发票数量", "签收数量", "数量"],
        # 用户 2026-09-19：模板最左新增「收货期间」列（决定这笔到票的月份归属）
        "receipt_period": ["收货期间", "收货月份", "入库期间", "入库月份"],
        "unit_price_excl_tax": ["不含税单价", "发票不含税单价", "单价"],
        "amount_excl_tax": ["不含税金额", "金额"],
    },
    "sales_invoice": {
        "invoice_no": ["发票号", "发票号码"],
        "invoice_date": ["发票日期", "开票日期", "日期"],
        "invoice_period": ["发票期间", "开票期间", "所属期间"],
        "sales_order_no": ["销售订单号", "订单号", "销售合同号"],
        "customer_code": ["客户编码", "客户代码"],
        "customer_name": ["客户", "客户名称"],
        "material_code": ["物料编码", "物料代码", "存货编码", "产品编码"],
        "material_name": ["物料名称", "存货名称", "产品名称", "品名"],
        "quantity": ["开票数量", "发票数量", "数量"],
        "unit_price_excl_tax": ["不含税单价", "销售单价", "单价"],
        "amount_excl_tax": ["不含税金额", "金额"],
    },
    "sales_ledger": {
        "out_month": ["出库期间"],
        "invoice_status": ["开票状态"],
        "invoice_month": ["开票期间"],
        "channel": ["销售渠道"],
        "company": ["销售公司"],
        "order_no": ["订单单号", "订单号", "销售订单号"],
        "customer_code": ["客户编码"],
        "customer_name": ["客户名称", "客户"],
        "sales_category": ["销售大类"],
        "material_code": ["产品编码", "物料编码", "物料编号", "存货编码", "产品编号"],
        "material_name": ["产品名称", "物料名称"],
        "specification": ["规格型号", "规格"],
        "uom": ["计量单位", "单位"],
        "terminal_amount": ["终端销售收入"],
        "terminal_unit_price": ["终端销售单价"],
        "lb_amount": ["射阳兰邦销售收入"],
        "lb_unit_price": ["射阳兰邦销售单价"],
        "quantity": ["销售数量", "数量"],
        "cost_price": ["销售成本单价"],
        "cost_amount": ["销售成本"],
        "gross_profit": ["销售毛利"],
        "gross_margin": ["销售毛利率"],
        "carry_type": ["结转类型"],
    },
    "work_order": {
        "order_no": ["工单编号", "生产工单号", "工单号"],
        "order_type": ["工单类型"],
        "identifier": ["标识码"],
        "plan_quantity": ["计划数量", "计划产量", "计划数"],
        "planned_start": ["计划开始时间", "计划开工时间", "计划开始日期"],
        "planned_finish": ["计划完工时间", "计划完成时间", "计划完工日期"],
        "priority": ["优先级"],
        "line_equipment": ["产线设备", "设备"],
        "shift": ["班次", "班别"],
        "production_dept_code": ["生产部门编号", "生产部门", "生产单位"],
        "production_supervisor_code": ["生产主管编号", "生产主管"],
        "planning_dept_code": ["计划部门编号", "计划部门"],
        "planner_code": ["计划员编号", "计划员"],
        "remark": ["备注"],
        "specified_material": ["指定用料"],
        "external_document_no": ["外部单据编号", "外部单号"],
        "status": ["状态", "工单状态"],
        "use_bom": ["是否使用物料清单", "使用物料清单"],
        "use_routing": ["是否使用工艺路线", "使用工艺路线"],
        "process_plan_type": ["工序计划类型"],
        "enable_sop": ["启用SOP", "是否启用SOP", "启用sop"],
        "production_mode": ["生产模式"],
        "line_code": ["产线编号", "生产线编号"],
        "source_order_no": ["订单号", "销售订单号", "关联订单号"],
        "material_color": ["物料颜色", "颜色"],
    },
    "material_issue": {
        "task_no": ["生产任务编号", "任务编号"],
        "operation_no": ["工序编号"],
        "operation_name": ["工序名称"],
        "material_code": ["物料编码", "物料编号", "物料代码"],
        "material_name": ["物料名称"],
        "material_category": ["物料分类", "分类"],
        "material_attribute": ["物料属性", "属性"],
        "material_spec": ["物料规格", "规格"],
        "issue_quantity": ["投料数量"],
        "returned_quantity": ["已回撤数量", "回撤数量"],
        "actual_issue_quantity": ["实际投料数量"],
        "issue_uom": ["投料单位", "单位"],
        "auxiliary_quantity": ["辅助单位数量", "辅助数量"],
        "auxiliary_uom": ["辅助单位"],
        "identifier": ["标识码"],
        "batch_no": ["批次号", "批号"],
        "quality_status": ["质量状态"],
        "issue_location": ["投料仓位", "仓位"],
        "replaced_material_code": ["被替代物料编号"],
        "replaced_material_name": ["被替代物料名称"],
        "issue_time": ["投料时间"],
        "issuer": ["投料人员", "投料人"],
        "executor": ["执行人员", "执行人"],
        "work_order_no": ["工单编号", "工单号", "生产工单号"],
        "main_output_code": ["主产出物料编号", "主产出物料编码"],
        "main_output_name": ["主产出物料名称"],
        "sales_order_no": ["销售订单编号", "销售订单号", "订单号"],
        "record_no": ["投料记录编号", "记录编号"],
    },
    "work_report": {
        "report_time": ["报工时间", "报工日期"],
        "task_no": ["生产任务编号", "任务编号"],
        "operation_no": ["工序编号"],
        "operation_name": ["工序名称"],
        "material_code": ["报工物料编号", "报工物料编码", "物料编码", "物料编号"],
        "material_name": ["报工物料名称", "物料名称"],
        "operator": ["生产人员", "操作人员", "报工人"],
        "quality_status": ["质量状态"],
        "defect_reason": ["不良原因"],
        "quantity": ["报工数量", "数量"],
        "uom": ["报工单位", "单位"],
        "auxiliary_quantity": ["辅助数量", "辅助单位数量"],
        "auxiliary_uom": ["辅助单位"],
        "batch_no": ["批次号", "批号"],
        "identifier": ["标识码"],
        "route_card_no": ["流转卡编号", "流转卡号"],
        "work_hours": ["工时"],
        "work_hours_uom": ["工时单位"],
        "remark": ["备注"],
        "department": ["部门"],
        "equipment": ["设备"],
        "work_order_no": ["工单编号", "工单号"],
        "work_order_material_code": ["工单物料编号", "工单物料编码"],
        "work_order_material_name": ["工单物料名称"],
        "order_no": ["订单编号", "订单号", "销售订单号"],
        "paper_tube_color": ["特纺/包覆-纸管颜色", "纸管颜色"],
        "spinning_shift": ["纺纱-班别", "班别"],
        "moisture_content": ["染色纤维-含水率", "含水率"],
        "dye_batch_flow": ["染色纤维-批号流水", "批号流水"],
        "weaving_sequence": ["织布-序号", "序号"],
        "weaving_machine": ["织布-机台", "机台"],
        "output_location": ["产出仓位", "产出库位"],
        "work_order_department": ["工单部门", "生产部门"],
    },
    "workshop_cost": {
        "workshop": ["车间", "生产车间", "成本中心"],
        "water_power": ["水电", "水电费", "电费"],
        "steam": ["蒸汽", "蒸汽费"],
        "natural_gas": ["天然气", "天然气费", "燃气费"],
        "depreciation": ["折旧", "折旧费"],
        "consumables": ["机物料", "机物料消耗", "辅料", "机物料费"],
        "dye_auxiliary": ["染料助剂", "染料助剂费", "助剂费"],
        "processing_fee": ["加工费", "加工费用"],
        "wages": ["工资", "人工", "直接人工", "人工费"],
    },
    "opening_wip": {
        "order_no": ["工单号", "生产订单号", "订单号"],
        "quantity": ["数量", "期初数量", "在产品数量", "期初在产品数量"],
        "material": ["材料", "材料成本", "期初材料"],
        "water_power": ["水电", "水电费"],
        "steam": ["蒸汽", "蒸汽费"],
        "natural_gas": ["天然气", "天然气费"],
        "depreciation": ["折旧", "折旧费"],
        "consumables": ["机物料", "机物料消耗"],
        "dye_auxiliary": ["染料助剂", "染料助剂费"],
        "processing_fee": ["加工费", "加工费用"],
        "wages": ["工资", "人工"],
    },
    "opening_purchase_accrual": {
        "original_period": ["原入库月份", "入库月份", "所属月份"],
        "document_date": ["原入库日期", "入库日期", "单据日期"],
        "document_no": ["原入库单号", "入库单号", "单据号"],
        "line_no": ["行号", "序号", "明细行"],
        "supplier_code": ["供应商编码", "供应商代码"],
        "supplier_name": ["供应商", "供应商名称"],
        "contract_no": ["合同号", "采购合同号"],
        "material_code": ["物料编码", "存货编码"],
        "material_name": ["物料名称", "存货名称"],
        "uom": ["单位", "计量单位"],
        "original_quantity": ["原入库数量", "累计入库数量"],
        "prior_invoiced_quantity": [
            "以前累计到票数量",
            "以前到票数量",
            "期前到票数量",
        ],
        "opening_quantity": [
            "月初未到票数量",
            "期初暂估数量",
            "期初未到票数量",
        ],
        "provisional_unit_price": [
            "暂估单价",
            "入库暂估单价",
            "不含税暂估单价",
        ],
        "opening_amount": [
            "月初暂估金额",
            "期初暂估金额",
            "未到票暂估金额",
        ],
        "remark": ["备注", "摘要"],
    },
    "opening_goods_issued": {
        # 与“销售台账（出库+开票）”完全一致的格式；仅用于系统首个核算月份之前的“已发货未开票”存量
        "row_no": ["序号"],
        "out_month": ["出库期间"],
        "invoice_status": ["开票状态"],
        "channel": ["销售渠道"],
        "company": ["销售公司"],
        "sales_order_no": ["订单单号", "订单号", "销售订单号"],
        "customer_code": ["客户编码", "客户编号"],
        "customer_name": ["客户名称", "客户"],
        "sales_category": ["销售大类"],
        "material_code": ["产品编码", "物料编码", "物料编号", "产品编号", "存货编码"],
        "material_name": ["产品名称", "物料名称"],
        "specification": ["规格型号", "规格"],
        "uom": ["计量单位", "单位"],
        "quantity": ["销售数量", "未开票数量", "数量"],
        "cost_price": ["销售成本单价", "原出库单位成本"],
        "cost_amount": ["销售成本", "月初发出商品成本"],
        "remark": ["备注", "摘要"],
    },
    "opening_inventory": {
        "material_code": ["物料号", "物料编码", "物料代码", "物料编号", "存货编码", "产品编码", "品号"],
        "material_name": ["物料名称", "存货名称", "产品名称", "品名"],
        "specification": ["规格", "规格型号", "型号"],
        "opening_date": ["日期", "期初日期", "结存日期", "业务日期"],
        "category_name": ["分类名称", "物料分类", "类别", "分类"],
        "warehouse": ["仓库名称", "仓库", "库位"],
        "batch_no": ["批次号", "批次", "批号"],
        "uom": ["单位名称", "单位", "计量单位", "基本单位"],
        "quantity": ["结存数量", "期初数量", "期初结存数量", "数量"],
        "unit_price": ["单价", "期初单价", "单位成本", "结存单价"],
        "amount": ["金额", "期初金额", "结存金额", "期初结存金额"],
        "remark": ["备注", "摘要"],
    },
    "ending_inventory": {
        "material_code": ["物料号", "物料编码", "物料代码", "物料编号", "存货编码", "产品编码", "品号"],
        "material_name": ["物料名称", "存货名称", "产品名称", "品名"],
        "specification": ["规格", "规格型号", "型号"],
        "ending_date": ["日期", "期末日期", "盘点日期", "结存日期"],
        "category_name": ["分类名称", "物料分类", "类别", "分类"],
        "warehouse": ["仓库名称", "仓库", "库位"],
        "batch_no": ["批次号", "批次", "批号"],
        "uom": ["单位名称", "单位", "计量单位", "基本单位"],
        "quantity": ["结存数量", "期末数量", "期末结存数量", "数量"],
        "unit_price": ["单价", "期末单价", "单位成本", "结存单价"],
        "amount": ["金额", "期末金额", "结存金额", "期末结存金额"],
        "remark": ["备注", "摘要"],
    },
}

# 浆粕专用资料类型：期初成品=期初存货格式，成品流水=库存流水格式，期末成品=期末存货格式
ALIASES["pulp_opening"] = ALIASES["opening_inventory"]
ALIASES["pulp_flow"] = ALIASES["inventory"]
ALIASES["pulp_ledger"] = ALIASES["inventory"]
ALIASES["pulp_ending"] = ALIASES["ending_inventory"]
ALIASES["pulp_material"] = {
    "txn_date": ["创建时间", "日期", "领料日期"],
    "material_code": ["物料编号", "物料编码", "物料代码", "编码"],
    "material_name": ["物料名称", "品名"],
    "category_name": ["分类名称", "分类"],
    "specification": ["物料规格", "规格"],
    "warehouse": ["仓库名称", "仓库"],
    "batch_no": ["批次号", "批号"],
    "quantity": ["操作数量", "数量", "领料数量"],
    "uom": ["单位名称", "单位", "计量单位"],
    "business_type": ["业务类型", "业务类别"],
    "workshop": ["车间", "车间名称"],
    "remark": ["备注", "摘要"],
    "unit_price": ["单价", "出库单价"],
    "amount": ["金额", "金额(元)"],
}
ALIASES["bom"] = {
    "work_order_no": ["工单编号", "工单号"],
    "material_code": ["物料编号", "成品物料编号", "父项物料编号"],
    "material_name": ["物料名称", "成品物料名称", "父项物料名称"],
    "line_no": ["项次", "行次"],
    "sub_item_code": ["子项物料编号", "子项物料编码", "子项编码"],
    "sub_item_name": ["子项物料名称", "子项名称"],
    "qty_numerator": ["子项用量:分子", "子项用量分子", "用量分子", "分子"],
    "qty_denominator": ["子项用量:分母", "子项用量分母", "用量分母", "分母"],
    "uom": ["单位名称", "单位", "计量单位"],
    "unit_price": ["单价（子项物料的单价）", "单价子项物料的单价", "子项单价", "单价"],
    "standard_price": ["标准成品单价", "标准单价", "成品标准单价"],
    "loss_rate": ["损耗率(%)", "损耗率", "损耗率%"],
    "standard_qty": ["标准用量"],
    "required_qty": ["需求数", "需求量", "BOM需求数", "需求数量"],
}

ALIASES["issue_progress"] = {
    "create_date": ["创建日期", "日期"],
    "department": ["部门名称", "部门", "车间"],
    "work_order_no": ["工单编号", "工单号"],
    "line_no": ["行号", "行次", "序号"],
    "material_code": ["物料编号", "物料编码", "物料代码"],
    "apply_quantity": ["申请数量", "申请量"],
    "issue_quantity": ["发料数量", "实发数量"],
    "issue_batch_mark": ["发料批次-标识码", "发料批次标识码", "发料批次", "标识码"],
    "actual_issue_quantity": ["投料数量", "实投数量"],
    "actual_batch_mark": ["投料批次-标识码", "投料批次标识码", "投料批次"],
    "uom": ["单位", "计量单位"],
    "issue_status": ["投料状态", "状态"],
}

ALIASES["historical_price"] = {
    "period": ["期间", "月份", "核算期间", "核算月份", "所属期间", "所属月份", "月度"],
    "material_code": ["物料编码", "物料编号", "物料代码", "产品编码", "存货编码", "编码"],
    "material_name": ["物料名称", "品名", "产品名称", "存货名称"],
    "receipt_price": [
        "入库单价", "入库价", "采购单价", "单价",
        "入库单价元", "入库价元", "采购单价元", "单价元",
    ],
    "issue_price": [
        "出库单价", "出库价", "销售单价", "发出单价",
        "出库单价元", "出库价元", "销售单价元", "发出单价元",
    ],
}

ALIASES["work_order_detail"] = {
    "created_at": ["创建时间", "创建日期", "日期"],
    "workshop": ["生产车间", "车间", "生产部门"],
    "order_no": ["工单编号", "工单号", "生产工单号"],
    "category": ["类别", "分类", "产品类别"],
    "planned_finish": ["计划完工时间", "计划完工", "计划完成时间"],
    "actual_finish": ["实际完工时间", "实际完工", "实际完成时间"],
    "completion_status": ["完成状态", "状态"],
    "plan_quantity": ["计划数量", "计划产量", "计划数"],
    "issue_quantity": ["投料数量", "投料量", "投料数"],
    "receipt_quantity": ["入库数量", "入库量", "完工入库数量", "报工数量"],
}


REQUIRED_FIELDS = {
    "inventory": {"txn_date", "business_type", "material_code", "quantity"},
    "pulp_opening": {"material_code", "quantity"},
    "pulp_flow": {"txn_date", "business_type", "material_code", "quantity"},
    "pulp_ending": {"material_code", "quantity"},
    "pulp_material": {"material_code", "quantity"},
    "pulp_ledger": {"txn_date", "business_type", "material_code", "quantity"},
    "bom": {"work_order_no", "sub_item_code"},
    "issue_progress": {"work_order_no", "material_code"},
    "historical_price": {"period", "material_code"},
    "work_order_detail": {"order_no"},
    "purchase_contract": {
        "contract_no",
        "material_code",
        "quantity",
        "unit_price_excl_tax",
    },
    "purchase_invoice": {
        "invoice_date",
        "material_code",
        "quantity",
        "unit_price_excl_tax",
    },
    "sales_invoice": {
        "invoice_no",
        "invoice_date",
        "material_code",
        "quantity",
    },
    "sales_ledger": {
        "out_month",
        "invoice_status",
        "material_code",
        "quantity",
    },
    "workshop_cost": {"workshop"},
    "opening_wip": {"order_no"},
    "opening_purchase_accrual": {
        "original_period",
        "document_no",
        "material_code",
        "original_quantity",
        "provisional_unit_price",
    },
    "opening_goods_issued": {
        "out_month",
        "invoice_status",
        "material_code",
        "quantity",
    },
    "opening_inventory": {
        "material_code",
        "quantity",
    },
    "ending_inventory": {
        "material_code",
        "quantity",
    },
    "work_order": {"order_no"},
    "material_issue": {"work_order_no", "material_code"},
    "work_report": {"work_order_no", "material_code", "quantity"},
}


def sha256_uploaded(uploaded_file) -> str:
    digest = hashlib.sha256()
    for chunk in uploaded_file.chunks():
        digest.update(chunk)
    uploaded_file.seek(0)
    return digest.hexdigest()


def to_decimal(value, default=None) -> Decimal | None:
    if value is None or value == "":
        return default
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return Decimal(int(value))
    text = str(value).strip().replace(",", "").replace("，", "").replace("￥", "").replace("¥", "")
    # 剥离常见中文单位后缀，如 88.88元/kg、3.36元/千克、5.5元
    text = re.sub(r"(元|円|美元|rmb)(/|／)?(kg|千克|公斤|吨|t|个|件|米|m|克|g)?$", "", text, flags=re.IGNORECASE).strip()
    if not text or text in {"-", "—", "－", "/", "\\", "None", "null", "n/a"}:
        return default
    is_percent = text.endswith("%")
    if is_percent:
        text = text[:-1]
    try:
        result = Decimal(text)
        return result / Decimal("100") if is_percent else result
    except InvalidOperation as exc:
        raise ValueError(f"无法识别数字：{value}") from exc


def to_date(value, default=None) -> date | None:
    if value is None or value == "":
        return default
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + __import__("datetime").timedelta(days=int(value))
    text = str(value).strip()
    if " " in text or "T" in text:
        text = text.split(" ")[0].split("T")[0]
    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y%m%d",
        "%Y-%m",
        "%Y/%m",
        "%Y%m",
    ):
        try:
            parsed = datetime.strptime(text, fmt).date()
            return parsed.replace(day=1) if "%d" not in fmt else parsed
        except ValueError:
            continue
    raise ValueError(f"无法识别日期：{value}")


def to_datetime(value, default=None) -> datetime | None:
    """支持日期、日期时间、Excel 序列号的通用时间解析。"""
    if value is None or value == "":
        return default
    if isinstance(value, datetime):
        return _aware_if_needed(value)
    if isinstance(value, date):
        return _aware_if_needed(datetime.combine(value, datetime.min.time()))
    if isinstance(value, (int, float)):
        return _aware_if_needed(
            datetime(1899, 12, 30) + __import__("datetime").timedelta(
                days=float(value)
            )
        )
    text = str(value).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y%m%d",
        "%Y%m%d%H%M%S",
    ):
        try:
            return _aware_if_needed(datetime.strptime(text, fmt))
        except ValueError:
            continue
    raise ValueError(f"无法识别时间：{value}")


def _aware_if_needed(value: datetime) -> datetime:
    if settings.USE_TZ and value.tzinfo is None:
        return dj_timezone.make_aware(value)
    return value


def to_bool(value, default=False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "y", "是", "有", "启用", "使用", "开"}:
        return True
    if text in {"0", "false", "no", "n", "否", "无", "不启用", "不使用", "关"}:
        return False
    return default


def month_start(value, fallback: date) -> date | None:
    parsed = to_date(value, fallback)
    if parsed is None:
        return None
    return parsed.replace(day=1)


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(path: Path):
    raw = path.read_bytes()
    decoded = None
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            decoded = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError("CSV文件编码无法识别，请使用UTF-8或GB18030。")
    return list(csv.reader(io.StringIO(decoded)))


def _read_xlsx(path: Path, kind: str = None):
    """读取 xlsx：kind 为 opening_wip 时返回所有有内容的 sheet 行列表（每车间一个 sheet）；
    其他类型返回最佳（内容最多）sheet 的行列表。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    best_rows = None
    best_score = -1
    sheet_rows_list = []
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        score = sum(1 for row in rows[:10] for value in row if value not in (None, ""))
        if score > best_score:
            best_score = score
            best_rows = rows
        if kind == "opening_wip" and any(
            value not in (None, "") for row in rows for value in row
        ):
            sheet_rows_list.append(rows)
    workbook.close()
    if kind == "opening_wip":
        return sheet_rows_list or ([best_rows] if best_rows else [])
    return [best_rows] if best_rows else []


def _locate_header(rows, kind: str):
    alias_map = {}
    for field, names in ALIASES[kind].items():
        for name in names:
            alias_map[normalize_header(name)] = field

    best = None
    for index, row in enumerate(rows[:10]):
        mapping = {}
        for col_index, value in enumerate(row):
            field = alias_map.get(normalize_header(value))
            if field and field not in mapping:
                mapping[field] = col_index
        required_hits = len(REQUIRED_FIELDS[kind] & set(mapping))
        score = required_hits * 100 + len(mapping)
        if best is None or score > best[0]:
            best = (score, index, mapping)

    if not best:
        raise ValueError("找不到表头。")
    _, header_index, mapping = best
    missing = REQUIRED_FIELDS[kind] - set(mapping)
    if missing:
        labels = "、".join(sorted(missing))
        raise ValueError(f"缺少必填字段：{labels}")
    return header_index, mapping


_PRICE_SHEET_IN = ("入库", "收货")
_PRICE_SHEET_OUT = ("出库", "发货")
_PRICE_CODE_HEADERS = {
    "品名编码", "物料编码", "物料编号", "产品编码", "存货编码", "产品编号",
}
_PRICE_NAME_HEADERS = {"品名名称", "物料名称", "产品名称"}
_MONTH_HEADER_RE = re.compile(r"(20\d{2})[^0-9]{0,3}(\d{1,2})")


def _wide_price_month(value):
    """表头「2026年7月 / 2026-07 / 202607」→ (年, 月)；解析不到返回 None。"""
    normalized = normalize_header(value)
    if not normalized:
        return None
    match = _MONTH_HEADER_RE.search(normalized)
    if not match:
        return None
    year, month = int(match.group(1)), int(match.group(2))
    if not 1 <= month <= 12:
        return None
    return year, month


def _read_wide_price_sheets(path: Path):
    """读取「单价表」双表格式（工作表名含 入库/出库，列=序号/品类/品名编码/品名名称 + 各月）。

    转成历史单价行；不是该格式时返回 None（由调用方按原格式处理）。
    只填了月份的物料才会出行；两个工作表同一物料同月会合并成一行。
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        book_sheets = []
        for worksheet in workbook.worksheets:
            title = worksheet.title or ""
            if any(token in title for token in _PRICE_SHEET_IN):
                side = "in"
            elif any(token in title for token in _PRICE_SHEET_OUT):
                side = "out"
            else:
                continue
            book_sheets.append(
                (side, [list(r) for r in worksheet.iter_rows(values_only=True)])
            )
    finally:
        workbook.close()
    if not book_sheets:
        return None

    existing_periods = set(
        AccountingPeriod.objects.values_list("month", flat=True)
    )
    merged = {}
    order = []
    skipped_months = set()
    for side, rows in book_sheets:
        header_index = None
        code_col = name_col = None
        month_cols = {}
        for index, row in enumerate(rows[:10]):
            local_code = local_name = None
            cols = {}
            for col_index, value in enumerate(row):
                normalized = normalize_header(value)
                if not normalized:
                    continue
                if local_code is None and normalized in _PRICE_CODE_HEADERS:
                    local_code = col_index
                elif local_name is None and normalized in _PRICE_NAME_HEADERS:
                    local_name = col_index
                else:
                    parsed = _wide_price_month(value)
                    if parsed is not None:
                        cols[col_index] = parsed
            if local_code is not None and cols:
                header_index = index
                code_col = local_code
                name_col = local_name
                month_cols = cols
                break
        if header_index is None:
            continue
        for row in rows[header_index + 1 :]:
            if not any(value not in (None, "") for value in row):
                continue
            code = text(row[code_col] if code_col < len(row) else None)
            if not code or code.startswith("示例"):
                continue
            name = (
                text(row[name_col])
                if name_col is not None and name_col < len(row)
                else ""
            )
            for col_index, (year, month) in sorted(month_cols.items()):
                value = row[col_index] if col_index < len(row) else None
                try:
                    price = to_decimal(value)
                except ValueError as exc:
                    raise ValueError(
                        f"{code} 的 {year}年{month}月 单价无法识别：{value!r}"
                    ) from exc
                if price is None or price == 0:
                    continue
                month_date = date(year, month, 1)
                if month_date not in existing_periods:
                    skipped_months.add(f"{year}年{month}月")
                    continue
                key = (month_date, code)
                item = merged.get(key)
                if item is None:
                    item = {
                        "_row_number": len(order) + 1,
                        "period": f"{year}-{month:02d}",
                        "material_code": code,
                        "material_name": name,
                        "_price_fields": set(),
                    }
                    merged[key] = item
                    order.append(key)
                if name and not item["material_name"]:
                    item["material_name"] = name
                if side == "in":
                    item["receipt_price"] = price
                    item["_price_fields"].add("receipt_price")
                else:
                    item["issue_price"] = price
                    item["_price_fields"].add("issue_price")

    emitted = [merged[key] for key in order]
    if not emitted:
        if skipped_months:
            raise ValueError(
                "这些月份还没有对应的核算期间，单价无法导入："
                + "、".join(sorted(skipped_months))
                + "。请先在系统里建立这些核算期间，或清除这些列的单价后重传。"
            )
        return None
    if skipped_months:
        emitted[0]["_note"] = (
            "⚠ 以下月份没有对应的核算期间，已跳过未导入："
            + "、".join(sorted(skipped_months))
        )
    return emitted


def read_rows(path: Path, kind: str):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows_list = [_read_csv(path)]
    elif suffix in {".xlsx", ".xlsm"}:
        if kind == "historical_price":
            wide_rows = _read_wide_price_sheets(path)
            if wide_rows is not None:
                return wide_rows
        rows_list = _read_xlsx(path, kind)
    else:
        raise ValueError("只支持CSV、XLSX和XLSM文件。")
    result = []
    for rows in rows_list:
        if not rows:
            continue
        if not any(value not in (None, "") for row in rows for value in row):
            continue
        header_index, mapping = _locate_header(rows, kind)
        for source_row_number, row in enumerate(
            rows[header_index + 1 :], header_index + 2
        ):
            if not any(value not in (None, "") for value in row):
                continue
            item = {"_row_number": source_row_number}
            for field, col_index in mapping.items():
                item[field] = row[col_index] if col_index < len(row) else None
            result.append(item)
    if not result:
        raise ValueError("文件没有可读取的数据。")
    return result


def infer_direction(value, business_type: str, quantity=None) -> str:
    normalized = text(value).lower()
    if normalized in {"入", "入库", "in", "收入"}:
        return InventoryTransaction.Direction.IN
    if normalized in {"出", "出库", "out", "发出"}:
        return InventoryTransaction.Direction.OUT

    # 明确方向的业务类型映射（先精确匹配，避免"报工调整"被"报工"误判）
    if text(business_type) == "报工":
        return InventoryTransaction.Direction.IN
    type_map = {
        "调拨发出": InventoryTransaction.Direction.OUT,
        "调拨接收": InventoryTransaction.Direction.IN,
        "投料撤回": InventoryTransaction.Direction.IN,
        "投料回撤": InventoryTransaction.Direction.IN,
        "投料": InventoryTransaction.Direction.OUT,
        "出库冲销": InventoryTransaction.Direction.IN,
        "冲销入库": InventoryTransaction.Direction.OUT,
        "退货入库": InventoryTransaction.Direction.IN,
        "采购退货": InventoryTransaction.Direction.OUT,
    }
    for token, direction in type_map.items():
        if token in business_type:
            return direction

    kind = business_type.lower()
    if "退料" in kind or "入库" in kind or "盘盈" in kind:
        return InventoryTransaction.Direction.IN
    if any(token in kind for token in ("出库", "领料", "发货", "盘亏", "采购退货")):
        return InventoryTransaction.Direction.OUT

    # 类型无法判断（库存调整、库存转换、报工调整、其它等）时，按数量正负判断
    if quantity is not None:
        try:
            q = Decimal(str(quantity).strip().replace(",", "") or "0")
            if q > 0:
                return InventoryTransaction.Direction.IN
            if q < 0:
                return InventoryTransaction.Direction.OUT
            # 数量为 0：无正负可判断，统一按入库处理（调用方会将 0 数量行标记为忽略）
            return InventoryTransaction.Direction.IN
        except InvalidOperation:
            pass
    raise ValueError(f"无法根据业务类型判断出入库方向：{business_type}")


# 车间/部门名称归一化：ERP 中文部门名 -> 部门编码
# （生产工单.生产部门编号 使用编码，如 tefang/JiZhi/RanZheng）
WORKSHOP_ALIASES = {
    "特纺": "tefang",
    "纺纱": "tefang",
    "机织": "JiZhi",
    "针织": "ZhenZhi",
    "包覆": "BaoFu",
    # 染色纤维 + 染色纱线 合并为 染色车间（产线/人员共用）
    "染色": "RanZheng",
    "染色车间": "RanZheng",
    "染色纤维": "RanZheng",
    "染整（纤维）": "RanZheng",
    "染整(纤维)": "RanZheng",
    "染整纤维": "RanZheng",
    "染色纱线": "RanZheng",
    "染整（纱线）": "RanZheng",
    "染整(纱线)": "RanZheng",
    "染整纱线": "RanZheng",
    # 旧编码（染色纱线/染色纤维车间已并入染色车间）也归一化到 RanZheng
    "RanZhengShaXian": "RanZheng",
    "RanZhengXianWei": "RanZheng",
    "纺纱车间": "tefang",
    "包覆车间": "BaoFu",
    "针织车间": "ZhenZhi",
    "机织车间": "JiZhi",
    "特纺投料": "tefang",
    "机织投料": "JiZhi",
    "针织投料": "ZhenZhi",
    "包覆投料": "BaoFu",
}

# 车间默认结转顺序（染色车间 → 纺纱 → 包覆 → 针织 → 机织；浆粕车间按需追加）
WORKSHOP_ORDER = [
    "RanZheng",
    "tefang",
    "BaoFu",
    "ZhenZhi",
    "JiZhi",
]

# 车间编码 -> 中文显示名
WORKSHOP_LABELS = {
    "RanZheng": "染色车间",
    "tefang": "纺纱车间",
    "BaoFu": "包覆车间",
    "ZhenZhi": "针织车间",
    "JiZhi": "机织车间",
    # 旧编码兜底显示（数据已归并，仅防历史残留显英文）
    "RanZhengShaXian": "染色车间",
    "RanZhengXianWei": "染色车间",
}


def normalize_workshop(value) -> str:
    """将 ERP 中文部门名归一化为部门编码；已是编码的原样返回。"""
    key = text(value)
    return WORKSHOP_ALIASES.get(key, key)


def get_or_create_material(code, name="", specification="", uom=""):
    material, _ = Material.objects.get_or_create(
        code=text(code),
        defaults={
            "name": text(name),
            "specification": text(specification),
            "uom": text(uom),
        },
    )
    changed = False
    if not material.name and name:
        material.name = text(name)
        changed = True
    if not material.specification and specification:
        material.specification = text(specification)
        changed = True
    if not material.uom and uom:
        material.uom = text(uom)
        changed = True
    if changed:
        material.save(update_fields=["name", "specification", "uom"])
    return material


def clear_calculation_outputs(period: AccountingPeriod):
    MatchAllocation.objects.filter(period=period).delete()
    InventoryValuation.objects.filter(period=period).delete()
    EndingInventoryBalance.objects.filter(
        period=period, is_system_generated=True
    ).delete()
    WorkOrderCostResult.objects.filter(period=period).delete()
    ReconciliationResult.objects.filter(period=period).delete()
    SettlementRun.objects.filter(period=period).delete()
    period.exceptions.all().delete()
    WorkshopSettlement.objects.filter(period=period).update(
        status=WorkshopSettlement.Status.PENDING,
        iterations=0,
        converged=False,
        settled_at=None,
        settled_by=None,
        order_cost=Decimal("0"),
        production_receipt_amount=Decimal("0"),
        error_message="",
    )
    period.status = AccountingPeriod.Status.OPEN
    period.last_run_at = None
    period.save(update_fields=["status", "last_run_at"])


def _delete_existing_kind(period: AccountingPeriod, kind: str):
    ImportBatch.objects.filter(period=period, kind=kind).delete()


def create_import_batch(*, period, kind, uploaded_file, user, replace_existing=False):
    if period.status == AccountingPeriod.Status.LOCKED:
        raise ValueError("该期间已锁账，不能继续导入。")
    if (
        kind
        in {
            ImportBatch.Kind.OPENING_PURCHASE_ACCRUAL,
            ImportBatch.Kind.OPENING_GOODS_ISSUED,
        }
        and AccountingPeriod.objects.filter(month__lt=period.month).exists()
    ):
        raise ValueError(
            "期初采购暂估和期初发出商品仅用于系统首个核算期间；"
            "后续月份由历史明细自动滚动。"
        )
    file_hash = sha256_uploaded(uploaded_file)
    if replace_existing:
        # 勾选"替换"时：先删除本期间同类旧批次，允许同名/同内容文件重新导入
        _delete_existing_kind(period, kind)
    elif ImportBatch.objects.filter(
        period=period, kind=kind, sha256=file_hash
    ).exists():
        raise ValueError("相同文件已经导入过，无需重复上传。")
    batch = ImportBatch.objects.create(
        period=period,
        kind=kind,
        file=uploaded_file,
        original_name=uploaded_file.name,
        sha256=file_hash,
        uploaded_by=user,
    )
    return batch


@transaction.atomic
def process_batch(batch: ImportBatch):
    rows = read_rows(Path(batch.file.path), batch.kind)
    # 历史单价是参考数据，不参与结算计算，导入时不清算本期结算输出
    if batch.kind not in {
        ImportBatch.Kind.HISTORICAL_PRICE,
        ImportBatch.Kind.WORK_ORDER_DETAIL,
    }:
        clear_calculation_outputs(batch.period)
    importer = {
        ImportBatch.Kind.INVENTORY: _import_inventory,
        ImportBatch.Kind.PURCHASE_CONTRACT: _import_purchase_contract,
        ImportBatch.Kind.PURCHASE_INVOICE: _import_purchase_invoice,
        ImportBatch.Kind.SALES_INVOICE: _import_sales_invoice,
        ImportBatch.Kind.SALES_LEDGER: _import_sales_ledger,
        ImportBatch.Kind.WORKSHOP_COST: _import_workshop_cost,
        ImportBatch.Kind.OPENING_WIP: _import_opening_wip,
        ImportBatch.Kind.OPENING_PURCHASE_ACCRUAL: (
            _import_opening_purchase_accrual
        ),
        ImportBatch.Kind.OPENING_GOODS_ISSUED: _import_opening_goods_issued,
        ImportBatch.Kind.OPENING_INVENTORY: _import_opening_inventory,
        ImportBatch.Kind.ENDING_INVENTORY: _import_ending_inventory,
        ImportBatch.Kind.WORK_ORDER: _import_work_order,
        ImportBatch.Kind.MATERIAL_ISSUE: _import_material_issue,
        ImportBatch.Kind.WORK_REPORT: _import_work_report,
        ImportBatch.Kind.PULP_OPENING: _import_opening_inventory,
        ImportBatch.Kind.PULP_FLOW: _import_inventory,
        ImportBatch.Kind.PULP_ENDING: _import_ending_inventory,
        ImportBatch.Kind.PULP_MATERIAL: _import_pulp_material,
        ImportBatch.Kind.PULP_LEDGER: _import_pulp_ledger,
        ImportBatch.Kind.BOM: _import_bom,
        ImportBatch.Kind.ISSUE_PROGRESS: _import_issue_progress,
        ImportBatch.Kind.HISTORICAL_PRICE: _import_historical_price,
        ImportBatch.Kind.WORK_ORDER_DETAIL: _import_work_order_detail,
    }[batch.kind]
    errors = []
    count = 0
    for row in rows:
        try:
            # 每行独立 savepoint：某行数据库错误只回滚该行，不拖垮整批
            with transaction.atomic():
                result = importer(batch, row)
            if result != "skipped":
                count += 1
        except Exception as exc:
            errors.append(f"第{row['_row_number']}行：{exc}")
    batch.row_count = count
    batch.error_count = len(errors)
    batch.message = "\n".join(errors[:100])
    notes = [str(row.get("_note")) for row in rows if row.get("_note")]
    if notes:
        batch.message = (
            (batch.message + "\n") if batch.message else ""
        ) + "\n".join(notes)
    if count == 0 and not errors:
        # 全部为模板示例行（跳过）或空数据：视为成功导入 0 行，不报错
        batch.status = ImportBatch.Status.SUCCESS
        batch.save(update_fields=["row_count", "error_count", "message", "status"])
        return batch
    batch.status = (
        ImportBatch.Status.SUCCESS if count > 0 else ImportBatch.Status.FAILED
    )
    batch.save(update_fields=["row_count", "error_count", "message", "status"])
    if count == 0:
        raise ValueError(batch.message or "没有成功导入任何数据。")
    if (
        batch.kind == ImportBatch.Kind.SALES_LEDGER
        and batch.status == ImportBatch.Status.SUCCESS
        and batch.error_count == 0
    ):
        # 台账整批成功后应用账务动作（生成开票记录/标记样品/核销），并把摘要附到消息
        from settlement.services.sales_ledger import apply_sales_ledger

        summary = apply_sales_ledger(batch)
        batch.message = ((batch.message + "\n") if batch.message else "") + summary
        batch.save(update_fields=["message"])
    return batch


def _import_inventory(batch, row):
    business_type = text(row.get("business_type"))
    document_no = text(row.get("document_no"))
    if not document_no and business_type == "报工":
        # 浆粕车间无报工记录，成品入库流水不填工单号时自动关联当月浆粕工单
        pulp_wo = WorkOrder.objects.filter(
            period=batch.period, production_dept_code="浆粕车间"
        ).first()
        if pulp_wo:
            document_no = pulp_wo.order_no
    raw_quantity = to_decimal(row.get("quantity"), Decimal("0"))
    direction = infer_direction(row.get("direction"), business_type, raw_quantity)
    quantity = abs(raw_quantity)
    is_zero_quantity = quantity == 0

    # 内部流转类型不参与成本核算：调拨、库存调整、库存转换、报工调整等；
    # 数量为 0 的行也作为忽略/无需匹配记录导入（行数与上传一致，不参与计价）。
    # 注意：投料/投料回撤/报工不再按内部流转忽略——
    #   投料出库按月末加权价出库、报工入库按工单完工成本入库，参与存货计价。
    INTERNAL_TRANSFER_TYPES = (
        "调拨",
        "库存调整",
        "库存转换",
        "报工调整",
    )
    is_internal = any(
        token in business_type for token in INTERNAL_TRANSFER_TYPES
    ) or is_zero_quantity

    material_code = text(row.get("material_code"))
    material = get_or_create_material(
        material_code,
        row.get("material_name"),
        row.get("specification"),
        row.get("uom"),
    )
    unit_price = to_decimal(row.get("unit_price"))
    amount = to_decimal(row.get("amount"))
    if amount is None and unit_price is not None:
        amount = (quantity * unit_price).quantize(Decimal("0.01"))
    InventoryTransaction.objects.create(
        period=batch.period,
        batch=batch,
        txn_date=to_date(row.get("txn_date")),
        document_no=document_no,
        line_no=text(row.get("line_no") or row["_row_number"]),
        business_type=business_type,
        direction=direction,
        warehouse=text(row.get("warehouse")),
        material=material,
        material_code=material_code,
        material_name=text(row.get("material_name")),
        specification=text(row.get("specification")),
        category_name=text(row.get("category_name")),
        batch_no=text(row.get("batch_no")),
        quantity=quantity,
        auxiliary_quantity=to_decimal(row.get("auxiliary_quantity")),
        uom=_clip_text(row.get("uom"), 30),
        auxiliary_uom=text(row.get("auxiliary_uom")),
        order_no=text(row.get("order_no")),
        workshop=text(row.get("workshop")),
        supplier_code=text(row.get("supplier_code")),
        supplier_name=text(row.get("supplier_name")),
        customer_code=text(row.get("customer_code")),
        customer_name=text(row.get("customer_name")),
        contract_no=text(row.get("contract_no")),
        invoice_no=text(row.get("invoice_no")),
        unit_price=unit_price,
        amount=amount,
        price_source="原始金额" if amount is not None else "",
        match_status=(
            InventoryTransaction.MatchStatus.NOT_REQUIRED
            if is_internal
            else InventoryTransaction.MatchStatus.UNMATCHED
        ),
        match_mode=(
            MatchingMode.IGNORED if is_internal else MatchingMode.AUTO
        ),
        remark=text(row.get("remark")),
    )


def _import_purchase_contract(batch, row):
    quantity = abs(to_decimal(row.get("quantity"), Decimal("0")))
    unit_price = to_decimal(row.get("unit_price_excl_tax"), Decimal("0"))
    # 「已签收数量」：有这一列就存下来（用户 2026-09-19：合同里有到货签收 → 暂估按合同收货期间归属）；
    # 文件里没有该列时存 None，含义是"未提供"，与"签收 0"区分开。
    received_raw = row.get("received_quantity")
    received = abs(to_decimal(received_raw, Decimal("0")))
    if unit_price > 0:
        # 合同数量取 max(合同数量, 已签收数量)：
        #  - 合同数量为空/0（如 R11.504）：用已签收数量兜底
        #  - 已签收超合同数量（超签，如 R82.203 5000/14650）：提升到已签收
        # 单价=0 的合同保持 0 数量，不参与匹配，流水继续用原始金额
        if received > quantity:
            quantity = received
    PurchaseContractLine.objects.create(
        period=batch.period,
        batch=batch,
        contract_no=text(row.get("contract_no")),
        contract_date=to_date(row.get("contract_date")),
        supplier_code=text(row.get("supplier_code")),
        supplier_name=text(row.get("supplier_name")),
        material_code=text(row.get("material_code")),
        material_name=text(row.get("material_name")),
        quantity=quantity,
        remaining_quantity=quantity,
        unit_price_excl_tax=unit_price,
        received_quantity=received if received_raw not in (None, "") else None,
        receipt_period=month_start(row.get("receipt_period"), None),
        # 收票日期（用户 2026-09-19）：留空＝尚未收票（仍挂暂估）；填了＝该批签收已到票
        invoice_date=to_date(row.get("invoice_date"), None),
        # 不含税总价（用户 2026-09-19）：文件给了就存下来，「合同金额」列以它为准
        amount_excl_tax=to_decimal(row.get("amount_excl_tax")),
    )


def _import_purchase_invoice(batch, row):
    quantity = abs(to_decimal(row.get("quantity"), Decimal("0")))
    unit_price = to_decimal(row.get("unit_price_excl_tax"), Decimal("0"))
    amount = to_decimal(row.get("amount_excl_tax"))
    if amount is None:
        amount = (quantity * unit_price).quantize(Decimal("0.01"))
    invoice_date = to_date(row.get("invoice_date"))
    if invoice_date is None:
        # 收票日期为空时，兜底用核算月份首日（发票期间随之落入核算月）
        invoice_date = batch.period.month
    invoice_no = text(row.get("invoice_no"))
    if not invoice_no:
        invoice_no = f"B{batch.id}-{row['_row_number']}"
    PurchaseInvoiceLine.objects.create(
        period=batch.period,
        batch=batch,
        invoice_no=invoice_no,
        invoice_date=invoice_date,
        invoice_period=month_start(row.get("invoice_period"), invoice_date),
        contract_no=text(row.get("contract_no")),
        supplier_code=text(row.get("supplier_code")),
        supplier_name=text(row.get("supplier_name")),
        material_code=text(row.get("material_code")),
        material_name=text(row.get("material_name")),
        uom=text(row.get("uom")),
        quantity=quantity,
        unit_price_excl_tax=unit_price,
        amount_excl_tax=amount,
        remaining_quantity=quantity,
        receipt_period=month_start(row.get("receipt_period"), None),
    )


def _import_sales_invoice(batch, row):
    quantity = abs(to_decimal(row.get("quantity"), Decimal("0")))
    unit_price = to_decimal(row.get("unit_price_excl_tax"))
    amount = to_decimal(row.get("amount_excl_tax"))
    if amount is None and unit_price is not None:
        amount = (quantity * unit_price).quantize(Decimal("0.01"))
    invoice_date = to_date(row.get("invoice_date"))
    if invoice_date is None:
        invoice_date = batch.period.month
    SalesInvoiceLine.objects.create(
        period=batch.period,
        batch=batch,
        invoice_no=text(row.get("invoice_no")),
        invoice_date=invoice_date,
        invoice_period=month_start(row.get("invoice_period"), invoice_date),
        sales_order_no=text(row.get("sales_order_no")),
        customer_code=text(row.get("customer_code")),
        customer_name=text(row.get("customer_name")),
        material_code=text(row.get("material_code")),
        material_name=text(row.get("material_name")),
        quantity=quantity,
        unit_price_excl_tax=unit_price,
        amount_excl_tax=amount,
        remaining_quantity=quantity,
    )


def _import_sales_ledger(batch, row):
    """销售台账（出库+开票）行导入：只校验并保存台账行，账务动作由 apply_sales_ledger 批量执行。"""
    # 模板示例行（订单单号或产品编码以“示例”开头）自动跳过，不导入也不计数
    if text(row.get("order_no")).startswith("示例") or text(
        row.get("material_code")
    ).startswith("示例"):
        return "skipped"
    status_text = text(row.get("invoice_status"))
    if not status_text:
        raise ValueError("开票状态不能为空（已开票/未开票/样品免费）。")
    if ("已开票" in status_text) or ("已开发票" in status_text):
        status = SalesLedgerLine.InvoiceStatus.INVOICED
    elif ("未开票" in status_text) or ("未开发票" in status_text):
        status = SalesLedgerLine.InvoiceStatus.UNINVOICED
    elif "样品" in status_text:
        status = SalesLedgerLine.InvoiceStatus.SAMPLE_FREE
    else:
        raise ValueError(
            f"开票状态无法识别：{status_text}"
            "（应为 已开票/未开票/样品免费，或 已开发票/未开发票）"
        )
    out_month = month_start(row.get("out_month"), None)
    if out_month is None:
        raise ValueError("出库期间不能为空或无法识别（格式如 2026-07）。")
    invoice_month = month_start(row.get("invoice_month"), None)
    if status == SalesLedgerLine.InvoiceStatus.INVOICED and invoice_month is None:
        raise ValueError("开票状态=已开票 时，开票期间必填（格式如 2026-07）。")
    if status != SalesLedgerLine.InvoiceStatus.INVOICED and invoice_month is not None:
        # 非“已开票”行多填了开票期间：以开票状态为准，忽略所填期间，仅在导入结果里提示
        row["_note"] = (
            f"第{row['_row_number']}行：开票状态={status_text}，"
            f"已忽略所填开票期间（{invoice_month:%Y-%m}）。"
        )
        invoice_month = None
    quantity = to_decimal(row.get("quantity"))
    if quantity is None or quantity <= 0:
        # 数量为空/0 的行没有金额可结转：跳过该行并提示，不影响其余行入账
        row["_note"] = (
            f"第{row['_row_number']}行：销售数量为空或0，已跳过（不参与结转）；"
            "如确属实际出货，请补数量后重新上传。"
        )
        return "skipped"
    code = text(row.get("material_code"))
    if not code:
        raise ValueError("产品编码不能为空。")
    material = get_or_create_material(
        code, row.get("material_name"), row.get("specification"), row.get("uom")
    )
    SalesLedgerLine.objects.create(
        period=batch.period,
        batch=batch,
        row_no=row["_row_number"],
        out_month=out_month,
        invoice_status=status,
        invoice_month=invoice_month,
        channel=text(row.get("channel"))[:120],
        company=text(row.get("company"))[:120],
        order_no=text(row.get("order_no")),
        customer_code=text(row.get("customer_code")),
        customer_name=text(row.get("customer_name")),
        sales_category=text(row.get("sales_category"))[:120],
        material=material,
        material_code=code,
        material_name=text(row.get("material_name"))[:255],
        specification=text(row.get("specification"))[:255],
        uom=text(row.get("uom"))[:30],
        terminal_amount=to_decimal(row.get("terminal_amount")),
        terminal_unit_price=to_decimal(row.get("terminal_unit_price")),
        lb_amount=to_decimal(row.get("lb_amount")),
        lb_unit_price=to_decimal(row.get("lb_unit_price")),
        quantity=quantity,
        cost_price=to_decimal(row.get("cost_price")),
        cost_amount=to_decimal(row.get("cost_amount")),
        gross_profit=to_decimal(row.get("gross_profit")),
        gross_margin=to_decimal(row.get("gross_margin")),
        carry_type=text(row.get("carry_type"))[:120],
    )


def _import_work_order(batch, row):
    order_no = text(row.get("order_no"))
    if not order_no:
        raise ValueError("工单编号不能为空。")
    WorkOrder.objects.update_or_create(
        period=batch.period,
        order_no=order_no,
        defaults={
            "batch": batch,
            "order_type": text(row.get("order_type")),
            "identifier": text(row.get("identifier")),
            "plan_quantity": abs(to_decimal(row.get("plan_quantity"), Decimal("0"))),
            "planned_start": to_datetime(row.get("planned_start")),
            "planned_finish": to_datetime(row.get("planned_finish")),
            "priority": text(row.get("priority")),
            "line_equipment": text(row.get("line_equipment")),
            "shift": text(row.get("shift")),
            "production_dept_code": normalize_workshop(text(row.get("production_dept_code"))),
            "production_supervisor_code": text(
                row.get("production_supervisor_code")
            ),
            "planning_dept_code": text(row.get("planning_dept_code")),
            "planner_code": text(row.get("planner_code")),
            "remark": text(row.get("remark")),
            "specified_material": text(row.get("specified_material")),
            "external_document_no": text(row.get("external_document_no")),
            "status": text(row.get("status")),
            "use_bom": to_bool(row.get("use_bom")),
            "use_routing": to_bool(row.get("use_routing")),
            "process_plan_type": text(row.get("process_plan_type")),
            "enable_sop": to_bool(row.get("enable_sop")),
            "production_mode": text(row.get("production_mode")),
            "line_code": text(row.get("line_code")),
            "source_order_no": text(row.get("source_order_no")),
            "material_color": text(row.get("material_color")),
        },
    )


def _import_material_issue(batch, row):
    work_order_no = _required_text(row.get("work_order_no"), 100, "工单编号")
    material_code = _required_text(row.get("material_code"), 80, "物料编号")
    if not work_order_no:
        raise ValueError("工单编号不能为空。")
    if not material_code:
        raise ValueError("物料编码不能为空。")
    issue_quantity = abs(to_decimal(row.get("issue_quantity"), Decimal("0")))
    returned_quantity = abs(
        to_decimal(row.get("returned_quantity"), Decimal("0"))
    )
    actual_issue_quantity = to_decimal(row.get("actual_issue_quantity"))
    if actual_issue_quantity is None:
        actual_issue_quantity = issue_quantity - returned_quantity
    actual_issue_quantity = abs(actual_issue_quantity)
    get_or_create_material(
        material_code,
        row.get("material_name"),
        row.get("material_spec"),
        row.get("issue_uom"),
    )
    MaterialIssue.objects.create(
        period=batch.period,
        batch=batch,
        task_no=text(row.get("task_no")),
        operation_no=text(row.get("operation_no")),
        operation_name=text(row.get("operation_name")),
        material_code=material_code,
        material_name=text(row.get("material_name")),
        material_category=text(row.get("material_category")),
        material_attribute=text(row.get("material_attribute")),
        material_spec=text(row.get("material_spec")),
        issue_quantity=issue_quantity,
        returned_quantity=returned_quantity,
        actual_issue_quantity=actual_issue_quantity,
        issue_uom=text(row.get("issue_uom")),
        auxiliary_quantity=to_decimal(row.get("auxiliary_quantity")),
        auxiliary_uom=text(row.get("auxiliary_uom")),
        identifier=text(row.get("identifier")),
        batch_no=text(row.get("batch_no")),
        quality_status=text(row.get("quality_status")),
        issue_location=text(row.get("issue_location")),
        replaced_material_code=text(row.get("replaced_material_code")),
        replaced_material_name=text(row.get("replaced_material_name")),
        issue_time=to_datetime(row.get("issue_time")),
        issuer=text(row.get("issuer")),
        executor=text(row.get("executor")),
        work_order_no=work_order_no,
        main_output_code=text(row.get("main_output_code")),
        main_output_name=text(row.get("main_output_name")),
        sales_order_no=text(row.get("sales_order_no")),
        record_no=text(row.get("record_no")),
    )


def _import_work_report(batch, row):
    work_order_no = text(row.get("work_order_no"))
    material_code = text(row.get("material_code"))
    if not work_order_no:
        raise ValueError("工单编号不能为空。")
    if not material_code:
        raise ValueError("报工物料编号不能为空。")
    quantity = abs(to_decimal(row.get("quantity"), Decimal("0")))
    if quantity == 0:
        raise ValueError("报工数量必须大于0。")
    get_or_create_material(
        material_code,
        row.get("material_name"),
        None,
        row.get("uom"),
    )
    WorkReport.objects.create(
        period=batch.period,
        batch=batch,
        report_time=to_datetime(row.get("report_time")),
        task_no=text(row.get("task_no")),
        operation_no=text(row.get("operation_no")),
        operation_name=text(row.get("operation_name")),
        material_code=material_code,
        material_name=text(row.get("material_name")),
        operator=text(row.get("operator")),
        quality_status=text(row.get("quality_status")),
        defect_reason=text(row.get("defect_reason")),
        quantity=quantity,
        uom=text(row.get("uom")),
        auxiliary_quantity=to_decimal(row.get("auxiliary_quantity")),
        auxiliary_uom=text(row.get("auxiliary_uom")),
        batch_no=text(row.get("batch_no")),
        identifier=text(row.get("identifier")),
        route_card_no=text(row.get("route_card_no")),
        work_hours=to_decimal(row.get("work_hours")),
        work_hours_uom=text(row.get("work_hours_uom")),
        remark=text(row.get("remark")),
        department=_clip_text(row.get("department"), 100),
        equipment=text(row.get("equipment")),
        work_order_no=work_order_no,
        work_order_material_code=text(row.get("work_order_material_code")),
        work_order_material_name=text(row.get("work_order_material_name")),
        order_no=text(row.get("order_no")),
        paper_tube_color=text(row.get("paper_tube_color")),
        spinning_shift=text(row.get("spinning_shift")),
        moisture_content=text(row.get("moisture_content")),
        dye_batch_flow=text(row.get("dye_batch_flow")),
        weaving_sequence=text(row.get("weaving_sequence")),
        weaving_machine=text(row.get("weaving_machine")),
        output_location=text(row.get("output_location")),
        work_order_department=text(row.get("work_order_department")),
    )


def _import_bom(batch, row):
    work_order_no = text(row.get("work_order_no"))
    sub_item_code = text(row.get("sub_item_code"))
    if not work_order_no:
        raise ValueError("工单编号不能为空。")
    if not sub_item_code:
        raise ValueError("子项物料编号不能为空。")
    denominator = to_decimal(row.get("qty_denominator"), Decimal("1"))
    if denominator is None or denominator <= 0:
        raise ValueError("子项用量分母必须大于0。")
    numerator = to_decimal(row.get("qty_numerator"), Decimal("0"))
    if numerator < 0:
        raise ValueError("子项用量分子不能小于0。")
    BomLine.objects.create(
        period=batch.period,
        batch=batch,
        work_order_no=work_order_no,
        material_code=text(row.get("material_code")),
        material_name=text(row.get("material_name")),
        line_no=_clip_text(row.get("line_no"), 40),
        sub_item_code=sub_item_code,
        sub_item_name=text(row.get("sub_item_name")),
        qty_numerator=numerator,
        qty_denominator=denominator,
        uom=text(row.get("uom")),
        unit_price=to_decimal(row.get("unit_price")),
        standard_price=to_decimal(row.get("standard_price")),
        loss_rate=to_decimal(row.get("loss_rate")),
        standard_qty=to_decimal(row.get("standard_qty")),
        required_qty=to_decimal(row.get("required_qty")),
    )


def _clip_text(value, max_len):
    """辅助字段超长时截断到 max_len（防 PostgreSQL value too long）。"""
    return text(value)[:max_len]


def _required_text(value, max_len, label):
    """关键字段超长时显式报错，使用户能看懂并修正。"""
    v = text(value)
    if len(v) > max_len:
        raise ValueError(f"{label}超长（最多 {max_len} 字符，当前 {len(v)} 个）。")
    return v


def _import_issue_progress(batch, row):
    work_order_no = _required_text(row.get("work_order_no"), 100, "工单编号")
    material_code = _required_text(row.get("material_code"), 80, "物料编号")
    if not work_order_no:
        raise ValueError("工单编号不能为空。")
    if not material_code:
        raise ValueError("物料编号不能为空。")
    IssueProgress.objects.create(
        period=batch.period,
        batch=batch,
        create_date=_clip_text(row.get("create_date"), 50),
        department=_clip_text(row.get("department"), 100),
        work_order_no=work_order_no,
        line_no=_clip_text(row.get("line_no"), 40),
        material_code=material_code,
        apply_quantity=to_decimal(row.get("apply_quantity"), Decimal("0")),
        issue_quantity=to_decimal(row.get("issue_quantity"), Decimal("0")),
        issue_batch_mark=_clip_text(row.get("issue_batch_mark"), 100),
        actual_issue_quantity=to_decimal(
            row.get("actual_issue_quantity"), Decimal("0")
        ),
        actual_batch_mark=_clip_text(row.get("actual_batch_mark"), 100),
        uom=_clip_text(row.get("uom"), 30),
        issue_status=_clip_text(row.get("issue_status"), 50),
    )


def _parse_period(value) -> AccountingPeriod:
    """解析期间列（2026-07 / 2026年07月 / 2026-7 / 202607 / 日期对象等）并返回对应 AccountingPeriod。"""
    if isinstance(value, (datetime, date)):
        return AccountingPeriod.objects.get_or_create(
            month=date(value.year, value.month, 1)
        )[0]
    raw = text(value)
    if not raw:
        raise ValueError("期间不能为空，格式如 2026-07。")
    m = re.search(r"(20\d{2})\D*(\d{1,2})", raw)
    if not m:
        raise ValueError(f"无法解析期间：{raw}，格式如 2026-07。")
    year, month = int(m.group(1)), int(m.group(2))
    if not 1 <= month <= 12:
        raise ValueError(f"月份无效：{month}。")
    period, _ = AccountingPeriod.objects.get_or_create(
        month=date(year, month, 1)
    )
    return period


def _import_historical_price(batch, row):
    """历史单价（月度）：模板示例行跳过。

    只更新本次填写的方向（入库/出库），未填的一侧保持原值不清空
    —— 避免"只填出库单价"的表格把已有入库单价抹掉。
    """
    material_code = text(row.get("material_code"))
    if material_code.startswith("示例"):
        return "skipped"
    if not material_code:
        raise ValueError("物料编码不能为空。")
    period = _parse_period(row.get("period"))
    receipt_price = to_decimal(row.get("receipt_price")) or None
    issue_price = to_decimal(row.get("issue_price")) or None
    provided = set(row.get("_price_fields") or ())
    if not provided:
        if receipt_price is not None:
            provided.add("receipt_price")
        if issue_price is not None:
            provided.add("issue_price")
    if not provided:
        raise ValueError("入库单价和出库单价至少填写一项。")
    target, _created = HistoricalPrice.objects.get_or_create(
        period=period, material_code=material_code
    )
    material_name = text(row.get("material_name"))
    if material_name:
        target.material_name = material_name
    if "receipt_price" in provided:
        target.receipt_price = receipt_price
    if "issue_price" in provided:
        target.issue_price = issue_price
    target.save()


def _import_work_order_detail(batch, row):
    order_no = text(row.get("order_no"))
    # 模板示例行（工单编号以"示例-"开头）自动跳过，不导入也不计数
    if order_no.startswith("示例-"):
        return "skipped"
    if not order_no:
        raise ValueError("工单编号不能为空。")
    WorkOrderDetail.objects.update_or_create(
        period=batch.period,
        order_no=order_no,
        defaults={
            "batch": batch,
            "created_at": text(row.get("created_at")),
            "workshop": text(row.get("workshop")),
            "category": text(row.get("category")),
            "planned_finish": text(row.get("planned_finish")),
            "actual_finish": text(row.get("actual_finish")),
            "completion_status": text(row.get("completion_status")),
            "plan_quantity": to_decimal(row.get("plan_quantity"), Decimal("0")),
            "issue_quantity": to_decimal(row.get("issue_quantity"), Decimal("0")),
            "receipt_quantity": to_decimal(row.get("receipt_quantity"), Decimal("0")),
            "row_no": int(row.get("_row_number") or 0),
        },
    )


def _import_workshop_cost(batch, row):
    workshop = normalize_workshop(text(row.get("workshop")))
    # 模板末行的"合计/总计"行仅作汇总提示，不导入
    if workshop in ("合计", "总计"):
        return
    WorkshopCostPool.objects.create(
        period=batch.period,
        batch=batch,
        workshop=workshop,
        water_power=to_decimal(row.get("water_power"), Decimal("0")),
        steam=to_decimal(row.get("steam"), Decimal("0")),
        natural_gas=to_decimal(row.get("natural_gas"), Decimal("0")),
        depreciation=to_decimal(row.get("depreciation"), Decimal("0")),
        consumables=to_decimal(row.get("consumables"), Decimal("0")),
        dye_auxiliary=to_decimal(row.get("dye_auxiliary"), Decimal("0")),
        wages=to_decimal(row.get("wages"), Decimal("0")),
    )


def _import_opening_wip(batch, row):
    order_no = text(row.get("order_no"))
    # 模板中的示例行（工单号以"示例-"开头）自动跳过，不导入也不计数
    if order_no.startswith("示例-"):
        return "skipped"
    OpeningWip.objects.update_or_create(
        period=batch.period,
        order_no=order_no,
        defaults={
            "batch": batch,
            "source_period": None,
            "is_system_generated": False,
            "quantity": to_decimal(row.get("quantity"), Decimal("0")),
            "material": to_decimal(row.get("material"), Decimal("0")),
            "water_power": to_decimal(row.get("water_power"), Decimal("0")),
            "steam": to_decimal(row.get("steam"), Decimal("0")),
            "natural_gas": to_decimal(
                row.get("natural_gas"), Decimal("0")
            ),
            "depreciation": to_decimal(
                row.get("depreciation"), Decimal("0")
            ),
            "consumables": to_decimal(
                row.get("consumables"), Decimal("0")
            ),
            "dye_auxiliary": to_decimal(
                row.get("dye_auxiliary"), Decimal("0")
            ),
            "processing_fee": to_decimal(row.get("processing_fee"), Decimal("0")),
            "wages": to_decimal(row.get("wages"), Decimal("0")),
        },
    )


def _opening_ledger_values(batch, row, *, price_field, ledger_name):
    original_period = month_start(row.get("original_period"), batch.period.month)
    if original_period >= batch.period.month:
        raise ValueError(f"{ledger_name}的原业务月份必须早于系统首个核算月份。")
    document_date = to_date(row.get("document_date"))
    if document_date is None:
        # 模板只填"原入库月份"时，默认取该月最后一天作为原单据日期
        import calendar

        last_day = calendar.monthrange(
            original_period.year, original_period.month
        )[1]
        document_date = original_period.replace(day=last_day)
    elif document_date >= batch.period.month:
        raise ValueError(f"{ledger_name}的原单据日期必须早于系统首个核算月份。")
    elif document_date.replace(day=1) != original_period:
        raise ValueError("原业务月份必须与原单据日期所在月份一致。")

    prior_quantity = to_decimal(
        row.get("prior_invoiced_quantity"), Decimal("0")
    )
    if prior_quantity is None or prior_quantity < 0:
        raise ValueError("以前累计核销数量不能小于0。")

    original_quantity = to_decimal(row.get("original_quantity"))
    opening_quantity = to_decimal(row.get("opening_quantity"))
    if original_quantity is None and opening_quantity is None:
        raise ValueError("原业务数量与期初未核销数量至少填一个。")
    if original_quantity is None:
        original_quantity = opening_quantity + prior_quantity
    if opening_quantity is None:
        opening_quantity = original_quantity - prior_quantity
    if original_quantity <= 0 or prior_quantity > original_quantity:
        raise ValueError("原业务数量或以前累计核销数量不合理。")
    if opening_quantity <= 0:
        raise ValueError("期初未核销数量必须大于0。")
    expected_opening = original_quantity - prior_quantity
    if abs(expected_opening - opening_quantity) > Decimal("0.000001"):
        raise ValueError(
            "期初未核销数量应等于原业务数量减以前累计核销数量。"
        )

    unit_price = to_decimal(row.get(price_field))
    if unit_price is None:
        raise ValueError(
            f"{ledger_name}的单价/单位成本未填写：请在模板的"
            f"“原出库单位成本”（或“暂估单价”）列填写该行金额，金额为0也请填0。"
        )
    if unit_price < 0:
        raise ValueError(f"{ledger_name}的单价/单位成本不能小于0。")
    opening_amount = to_decimal(row.get("opening_amount"))
    if opening_amount is None:
        opening_amount = (opening_quantity * unit_price).quantize(
            Decimal("0.01")
        )
    if opening_amount < 0:
        raise ValueError("期初金额不能小于0。")

    get_or_create_material(
        row.get("material_code"),
        row.get("material_name"),
        uom=row.get("uom"),
    )
    return {
        "original_period": original_period,
        "document_date": document_date,
        "document_no": text(row.get("document_no")),
        "line_no": text(row.get("line_no")) or str(row["_row_number"]),
        "material_code": text(row.get("material_code")),
        "material_name": text(row.get("material_name")),
        "uom": text(row.get("uom")),
        "original_quantity": original_quantity,
        "prior_invoiced_quantity": prior_quantity,
        "opening_quantity": opening_quantity,
        "opening_amount": opening_amount,
        "remaining_quantity": opening_quantity,
        "remark": text(row.get("remark")),
    }, unit_price


def _import_opening_purchase_accrual(batch, row):
    values, unit_price = _opening_ledger_values(
        batch,
        row,
        price_field="provisional_unit_price",
        ledger_name="期初采购暂估",
    )
    OpeningPurchaseAccrual.objects.create(
        period=batch.period,
        batch=batch,
        supplier_code=text(row.get("supplier_code")),
        supplier_name=text(row.get("supplier_name")),
        contract_no=text(row.get("contract_no")),
        provisional_unit_price=unit_price,
        **values,
    )


def _import_opening_goods_issued(batch, row):
    """期初发出商品明细（与销售台账同格式）：首期前的“已发货未开票”存量。

    与销售台账的差异：
    - 只接受 开票状态=未开票 的行（已开票/样品免费请放在对应月份的销售台账里结转）；
    - 出库期间必须早于系统首个核算月份；
    - 销售成本单价/销售成本 必填其一——期初行没有对应库存流水，成本无法按月末加权算，
      以你上传的历史成本为准（当月销售台账才允许成本留空）。
    """
    if text(row.get("sales_order_no")).startswith("示例") or text(
        row.get("material_code")
    ).startswith("示例"):
        return "skipped"
    status_text = text(row.get("invoice_status"))
    if status_text and "未开票" not in status_text:
        raise ValueError(
            f"期初发出商品明细只能填“未开票”行，当前为：{status_text}；"
            "已开票/样品免费的行请放到对应出库月份的销售台账里。"
        )
    if text(row.get("invoice_month")):
        raise ValueError("期初发出商品明细行不应填写开票期间（全部为已发货未开票）。")
    out_month = month_start(row.get("out_month"), None)
    if out_month is None:
        raise ValueError("出库期间不能为空或无法识别（格式如 2026-06）。")
    if out_month >= batch.period.month:
        raise ValueError(
            "期初发出商品明细的出库期间必须早于系统首个核算月份（后续月份由台账自动滚动）。"
        )
    code = text(row.get("material_code"))
    if not code:
        raise ValueError("产品编码不能为空。")
    quantity = to_decimal(row.get("quantity"))
    if quantity is None or quantity <= 0:
        raise ValueError("销售数量（未开票数量）必须大于0。")
    get_or_create_material(
        code, row.get("material_name"), row.get("specification"), row.get("uom")
    )
    cost_price = to_decimal(row.get("cost_price"))
    cost_amount = to_decimal(row.get("cost_amount"))
    if cost_price is None and cost_amount is None:
        raise ValueError(
            "期初发出商品明细需填 销售成本单价 或 销售成本（金额）：期初为系统首月前的历史发货，"
            "没有流水可做月末加权，成本以你上传的历史成本为准（仅当月销售台账可留空成本）。"
        )
    if cost_amount is None:
        cost_amount = (quantity * cost_price).quantize(Decimal("0.01"))
    if cost_price is None:
        cost_price = (cost_amount / quantity).quantize(Decimal("0.00000001"))
    if cost_price < 0 or cost_amount < 0:
        raise ValueError("销售成本单价/销售成本不能小于0。")
    import calendar

    last_day = calendar.monthrange(out_month.year, out_month.month)[1]
    seq_no = text(row.get("row_no")) or str(row["_row_number"])
    document_no = text(row.get("sales_order_no")) or (
        f"OG-{batch.period.month:%Y%m}-{row['_row_number']}"
    )
    OpeningGoodsIssued.objects.create(
        period=batch.period,
        batch=batch,
        original_period=out_month,
        document_date=out_month.replace(day=last_day),
        document_no=document_no[:100],
        line_no=seq_no[:40],
        sales_order_no=text(row.get("sales_order_no")),
        customer_code=text(row.get("customer_code")),
        customer_name=text(row.get("customer_name")),
        sales_category=(text(row.get("sales_category")) or "")[:120],
        material_code=code,
        material_name=text(row.get("material_name"))[:255],
        uom=text(row.get("uom"))[:30],
        original_quantity=quantity,
        prior_invoiced_quantity=Decimal("0"),
        opening_quantity=quantity,
        unit_cost=cost_price,
        opening_amount=cost_amount,
        remaining_quantity=quantity,
        remark=text(row.get("remark"))[:500],
    )
def _import_opening_inventory(batch, row):
    material_code = text(row.get("material_code"))
    if not material_code:
        raise ValueError("物料号不能为空。")
    quantity = to_decimal(row.get("quantity"))
    if quantity is None or quantity <= 0:
        raise ValueError("结存数量必须大于0。")
    unit_price = to_decimal(row.get("unit_price"))
    amount = to_decimal(row.get("amount"))
    if amount is None and unit_price:
        amount = (quantity * unit_price).quantize(Decimal("0.01"))
    if not unit_price and amount:
        unit_price = (amount / quantity).quantize(Decimal("0.00000001"))
    # 单价/金额可都为空：数量以用户上传为准，单价金额由结算时回填
    unit_price = unit_price or Decimal("0")
    amount = amount or Decimal("0")
    if amount < 0 or unit_price < 0:
        raise ValueError("单价和金额不能小于0。")
    get_or_create_material(
        material_code,
        row.get("material_name"),
        row.get("specification"),
        row.get("uom"),
    )
    # 按 物料+仓库+批次 维度存储，保留各仓库/批次明细行；同维度累加
    warehouse = text(row.get("warehouse")) or None
    batch_no = text(row.get("batch_no")) or None
    item, created = OpeningInventoryBalance.objects.get_or_create(
        period=batch.period,
        material_code=material_code,
        warehouse=warehouse,
        batch_no=batch_no,
        defaults={
            "batch": batch,
            "source_period": None,
            "material_name": text(row.get("material_name")),
            "opening_date": to_date(row.get("opening_date")),
            "category_name": text(row.get("category_name")),
            "uom": text(row.get("uom")),
            "quantity": quantity,
            "amount": amount,
            "unit_price": unit_price or Decimal("0"),
            "is_system_generated": False,
        },
    )
    if not created:
        item.quantity = (item.quantity or Decimal("0")) + quantity
        item.amount = (item.amount or Decimal("0")) + amount
        if item.quantity:
            item.unit_price = (item.amount / item.quantity).quantize(
                Decimal("0.00000001")
            )
        if not item.material_name:
            item.material_name = text(row.get("material_name"))
        if not item.category_name:
            item.category_name = text(row.get("category_name"))
        if not item.uom:
            item.uom = text(row.get("uom"))
        item.batch = batch
        item.is_system_generated = False
        item.save()


def _import_ending_inventory(batch, row):
    material_code = text(row.get("material_code"))
    if not material_code:
        raise ValueError("物料号不能为空。")
    quantity = to_decimal(row.get("quantity"))
    if quantity is None or quantity <= 0:
        raise ValueError("结存数量必须大于0。")
    unit_price = to_decimal(row.get("unit_price"))
    amount = to_decimal(row.get("amount"))
    if amount is None and unit_price:
        amount = (quantity * unit_price).quantize(Decimal("0.01"))
    if not unit_price and amount:
        unit_price = (amount / quantity).quantize(Decimal("0.00000001"))
    # 单价/金额可都为空：数量以用户上传为准，单价金额由结算时回填
    unit_price = unit_price or Decimal("0")
    amount = amount or Decimal("0")
    if amount < 0 or unit_price < 0:
        raise ValueError("单价和金额不能小于0。")
    get_or_create_material(
        material_code,
        row.get("material_name"),
        row.get("specification"),
        row.get("uom"),
    )
    # 按 物料+仓库+批次 维度存储，保留各仓库/批次明细行；同维度累加
    warehouse = text(row.get("warehouse")) or None
    batch_no = text(row.get("batch_no")) or None
    item, created = EndingInventoryBalance.objects.get_or_create(
        period=batch.period,
        material_code=material_code,
        warehouse=warehouse,
        batch_no=batch_no,
        defaults={
            "batch": batch,
            "material_name": text(row.get("material_name")),
            "ending_date": to_date(row.get("ending_date")),
            "category_name": text(row.get("category_name")),
            "uom": text(row.get("uom")),
            "quantity": quantity,
            "amount": amount,
            "unit_price": unit_price or Decimal("0"),
            "is_system_generated": False,
        },
    )
    if not created:
        item.quantity = (item.quantity or Decimal("0")) + quantity
        item.amount = (item.amount or Decimal("0")) + amount
        if item.quantity:
            item.unit_price = (item.amount / item.quantity).quantize(
                Decimal("0.00000001")
            )
        if not item.material_name:
            item.material_name = text(row.get("material_name"))
        if not item.category_name:
            item.category_name = text(row.get("category_name"))
        if not item.uom:
            item.uom = text(row.get("uom"))
        item.batch = batch
        item.is_system_generated = False
        item.save()


def _import_pulp_material(batch, row):
    """浆粕原料流水（用户格式：创建时间/物料编号/物料名称/分类名称/物料规格/仓库名称/批次号/
    操作数量/单位名称/业务类型/车间/备注）。材料成本只认业务类型为领料出库的行；
    单价/金额可选：有则用用户值，都没有时金额由结算按历史单价（出库优先、入库兜底）计算。"""
    material_code = text(row.get("material_code"))
    material_name = text(row.get("material_name"))
    if (
        (not material_code)
        or material_code.startswith("示例")
        or material_name.startswith("示例")
    ):
        return "skipped"
    quantity = abs(to_decimal(row.get("quantity"), Decimal("0")))
    unit_price = to_decimal(row.get("unit_price"))
    amount = to_decimal(row.get("amount"))
    if amount is None and unit_price is not None:
        amount = (quantity * unit_price).quantize(Decimal("0.01"))
    PulpMaterial.objects.create(
        period=batch.period,
        batch=batch,
        txn_date=to_date(row.get("txn_date")),
        document_no=text(row.get("document_no")),
        material_code=material_code,
        material_name=material_name,
        category_name=text(row.get("category_name")),
        specification=text(row.get("specification")),
        warehouse=text(row.get("warehouse")),
        batch_no=text(row.get("batch_no")),
        quantity=quantity,
        uom=text(row.get("uom")),
        business_type=text(row.get("business_type")),
        workshop=text(row.get("workshop")),
        unit_price=unit_price,
        amount=amount if amount is not None else Decimal("0"),
        remark=text(row.get("remark")),
    )


def _import_pulp_ledger(batch, row):
    """浆粕流水（原料+成品通用模板）：按业务类型分流。
    领料出库 → 浆粕原料流水（材料成本来源）；生产入库/销售出库 → 库存流水（成品）。"""
    material_code = text(row.get("material_code"))
    material_name = text(row.get("material_name"))
    if (
        (not material_code)
        or material_code.startswith("示例")
        or material_name.startswith("示例")
    ):
        return "skipped"
    business_type = text(row.get("business_type"))
    if not business_type:
        return "skipped"  # 空行/说明行（无业务类型）自动跳过
    if "领料出库" in business_type:
        return _import_pulp_material(batch, row)
    if business_type in ("生产入库", "销售出库"):
        return _import_inventory(batch, row)
    raise ValueError(
        f"浆粕流水业务类型「{business_type}」不支持：请用 领料出库/生产入库/销售出库。"
    )
